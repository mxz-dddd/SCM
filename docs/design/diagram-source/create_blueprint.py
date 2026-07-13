from __future__ import annotations

import json
import math
import os
import shutil
import textwrap
import zipfile
from pathlib import Path
from typing import Dict, List, Tuple, Optional, Iterable

from graphviz import Source
import matplotlib.pyplot as plt
from matplotlib import font_manager
from PIL import Image

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path('/mnt/data/scm_blueprint')
DIAG = BASE / 'diagrams'
DIAG.mkdir(parents=True, exist_ok=True)

FONT_NAME = 'Noto Sans CJK SC'
FONT_PATH = '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
FONT_BOLD_PATH = '/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc'
try:
    font_manager.fontManager.addfont(FONT_PATH)
    font_manager.fontManager.addfont(FONT_BOLD_PATH)
except Exception:
    pass
plt.rcParams['axes.unicode_minus'] = False

PALETTE = {
    'navy': '#16324F', 'blue': '#2A6FDB', 'sky': '#EAF2FF',
    'cyan': '#DDF7F4', 'teal': '#087F8C', 'green': '#E8F5E9',
    'amber': '#FFF4D6', 'orange': '#C56A00', 'rose': '#FDECEC',
    'red': '#B42318', 'gray': '#F3F5F7', 'gray2': '#D6DCE4',
    'text': '#1F2937', 'muted': '#5B6472', 'white': '#FFFFFF',
    'violet': '#E8EDFF', 'purple': '#596BB3'
}


def render_dot(name: str, dot: str, engine: str = 'dot') -> Path:
    dot_path = DIAG / f'{name}.dot'
    dot_path.write_text(dot, encoding='utf-8')
    for fmt in ('svg', 'png'):
        src = Source(dot, engine=engine, filename=name, directory=str(DIAG), format=fmt)
        src.render(cleanup=True)
    return DIAG / f'{name}.png'


def graph_header(rankdir='TB', nodesep=0.32, ranksep=0.55, splines='spline') -> str:
    return f'''digraph G {{
      graph [rankdir={rankdir}, bgcolor="white", pad="0.18", nodesep="{nodesep}", ranksep="{ranksep}", splines={splines}, fontname="{FONT_NAME}", labelloc="t"];
      node [shape=box, style="rounded,filled", fontname="{FONT_NAME}", fontsize=10.5, color="{PALETTE['gray2']}", fillcolor="white", fontcolor="{PALETTE['text']}", margin="0.12,0.08"];
      edge [fontname="{FONT_NAME}", fontsize=8.5, color="#6B7280", fontcolor="#374151", arrowsize=0.7];
'''


def q(s: str) -> str:
    return s.replace('"', '\\"')


def architecture_diagrams() -> Dict[str, Path]:
    out: Dict[str, Path] = {}

    dot = graph_header('LR', 0.45, 0.85, 'spline')
    dot += f'''
      internal [label="内部用户\\n订单客服 / 仓库 / 运输 / 财务 / 管理层", fillcolor="{PALETTE['sky']}", color="{PALETTE['blue']}"];
      partner [label="外部协同方\\n客户 / 供应商 / 承运商 / 司机", fillcolor="{PALETTE['cyan']}", color="{PALETTE['teal']}"];
      platform [label="供应链协同云平台\\nOMS · WMS · TMS · AMS · 结算 · 控制塔", fillcolor="#DCE9FF", color="{PALETTE['blue']}", penwidth=2.1, fontsize=13, width=3.2, height=1.0];
      enterprise [label="企业系统\\nERP / MES / 电商 / 财务 / 主数据", fillcolor="{PALETTE['gray']}"];
      device [label="现场设备与 IoT\\nRF / 打印 / 称重 / 门禁 / GPS / 温控", fillcolor="{PALETTE['amber']}", color="#C58B00"];
      service [label="外部服务\\n地图 / 短信 / 邮件 / 微信 / 推送", fillcolor="{PALETTE['gray']}"];
      internal -> platform [label="Web / RF / APP"];
      partner -> platform [label="门户 / APP / 小程序"];
      enterprise -> platform [dir=both, label="API / EDI / 文件"];
      device -> platform [dir=both, label="遥测 / 指令"];
      service -> platform [dir=both, label="通知 / 地图 / 轨迹"];
    }}'''
    out['01_system_context'] = render_dot('01_system_context', dot)

    dot = graph_header('TB', 0.35, 0.62, 'spline')
    dot += f'''
      root [label="供应链协同云能力地图", fillcolor="{PALETTE['navy']}", fontcolor="white", color="{PALETTE['navy']}", fontsize=14, width=4.2, height=0.7];
      core [label="核心业务域\\nOMS：订单、分配、协同、履约\\nWMS：入库、库存、出库、增值、劳务\\nTMS：计划、配载、委托、跟踪、回单\\nAMS：预约、排队、门岗、月台\\nBilling：计费、凭证、对账、分摊", fillcolor="{PALETTE['sky']}", color="{PALETTE['blue']}", width=7.2, height=2.15, fontsize=11];
      common [label="平台与共性能力\\n租户 / 组织 / 用户 / 角色 / 数据权限\\n主数据 / 字典 / 合同 / 费率 / 地址\\n工作流 / 规则 / 策略 / 调度任务\\n附件 / 标签 / 打印 / 评论 / 审计\\nOpen API / EDI / Webhook / 适配器", fillcolor="{PALETTE['green']}", color="{PALETTE['teal']}", width=7.2, height=2.15, fontsize=11];
      insight [label="可视化、协同与优化\\n控制塔 / 事件 / 例外 / SLA\\nBI / KPI / 报表 / 大屏\\n路径 / 装载 / 需求 / 库存 / 网络优化\\n客户 / 司机 / 仓库 / 合作伙伴移动端", fillcolor="{PALETTE['violet']}", color="{PALETTE['purple']}", width=7.2, height=1.9, fontsize=11];
      root -> core;
      core -> common [label="由平台能力支撑", dir=back];
      core -> insight [label="业务事实与指标"];
      common -> insight [label="事件、查询、通知"];
    }}'''
    out['02_capability_map'] = render_dot('02_capability_map', dot)

    dot = graph_header('TB', 0.20, 0.38, 'spline')
    dot += f'''
      subgraph cluster_channel {{ label="体验与渠道层"; color="{PALETTE['gray2']}"; style="rounded";
        web [label="管理工作台 Web", fillcolor="{PALETTE['sky']}"];
        rf [label="WMS RF / PDA", fillcolor="{PALETTE['sky']}"];
        apps [label="司机 / 承运商 / 客户 APP", fillcolor="{PALETTE['sky']}"];
        portal [label="伙伴门户 / 小程序", fillcolor="{PALETTE['sky']}"];
        screen [label="控制塔大屏 / BI", fillcolor="{PALETTE['sky']}"];
      }}
      access [label="CDN / WAF / API Gateway / SSO / Tenant Resolver / Rate Limit", fillcolor="#DCE9FF", color="{PALETTE['blue']}", penwidth=1.5];
      subgraph cluster_app {{ label="应用编排层"; color="{PALETTE['gray2']}"; style="rounded";
        workbff [label="Workbench BFF"]; mobilebff [label="Mobile BFF"]; query [label="Query / Search API"]; realtime [label="Realtime Gateway"];
      }}
      subgraph cluster_domain {{ label="领域服务层（逻辑服务化；首期可模块化单体）"; color="{PALETTE['blue']}"; style="rounded"; penwidth=1.4;
        iam [label="IAM / Tenant / Org"]; mdm [label="MDM / Config"]; oms [label="OMS"]; wms [label="WMS"]; tms [label="TMS"]; ams [label="AMS"]; bill [label="Billing / Settlement"]; ctl [label="Event / Exception / Control Tower"]; doc [label="Document / Print / POD"]; notify [label="Notification"]; integ [label="Integration Hub"]; opt [label="Optimization / AI"];
      }}
      subgraph cluster_platform {{ label="平台基础能力"; color="{PALETTE['gray2']}"; style="rounded";
        flow [label="Workflow / Rule Engine"]; job [label="Scheduler / Async Job"]; bus [label="Event Bus / Outbox / Inbox", fillcolor="{PALETTE['amber']}"]; obs [label="Audit / Logs / Metrics / Traces"]; flag [label="Feature Flag / Tenant Config"];
      }}
      subgraph cluster_data {{ label="数据与基础设施"; color="{PALETTE['gray2']}"; style="rounded";
        db [label="事务数据库（按域 / Schema 隔离）"]; redis [label="Redis Cache / Lock"]; obj [label="Object Storage"]; search [label="Search / GIS / Time-series"]; lake [label="CDC / Data Lake / OLAP"];
      }}
      web -> access; rf -> access; apps -> access; portal -> access; screen -> access;
      access -> workbff; access -> mobilebff; access -> query; access -> realtime;
      workbff -> oms; workbff -> wms; workbff -> tms; workbff -> ams; workbff -> bill; workbff -> ctl;
      mobilebff -> wms; mobilebff -> tms; mobilebff -> ams;
      query -> search; realtime -> bus;
      iam -> db; mdm -> db; oms -> db; wms -> db; tms -> db; ams -> db; bill -> db;
      oms -> bus; wms -> bus; tms -> bus; ams -> bus; bill -> bus; ctl -> bus;
      flow -> oms; flow -> wms; flow -> tms; flow -> ams; flow -> bill;
      integ -> bus; doc -> obj; notify -> bus; opt -> bus; bus -> lake; db -> lake [label="CDC"];
      obs -> oms [style=dashed, constraint=false]; obs -> wms [style=dashed, constraint=false]; obs -> tms [style=dashed, constraint=false];
    }}'''
    out['03_logical_architecture'] = render_dot('03_logical_architecture', dot)

    dot = graph_header('TB', 0.28, 0.48, 'spline')
    dot += f'''
      internet [label="Internet / 企业专线 / 移动网络", shape=oval, fillcolor="{PALETTE['gray']}"];
      subgraph cluster_edge {{ label="边缘与接入区"; color="{PALETTE['gray2']}"; style="rounded";
        waf [label="WAF / CDN / DDoS"]; lb [label="Load Balancer / API Gateway"]; idp [label="Identity Provider / SSO"];
      }}
      subgraph cluster_app {{ label="应用集群（无状态、水平扩展）"; color="{PALETTE['blue']}"; style="rounded";
        webpod [label="Web / BFF Pods × N", fillcolor="{PALETTE['sky']}"]; domain [label="Domain Service Pods × N", fillcolor="{PALETTE['sky']}"]; worker [label="Async Worker / Scheduler × N", fillcolor="{PALETTE['sky']}"]; realtime [label="Realtime / Push × N", fillcolor="{PALETTE['sky']}"];
      }}
      subgraph cluster_data {{ label="数据区（多可用区）"; color="{PALETTE['teal']}"; style="rounded";
        dbp [label="Primary DB / HA"]; dbr [label="Read Replica / Reporting"]; redis [label="Redis Cluster"]; mq [label="Message Broker Cluster"]; storage [label="Object Storage / Backup"]; search [label="Search / GIS / Time-series"]; olap [label="OLAP / Data Lake"];
      }}
      subgraph cluster_ops {{ label="运维与安全"; color="{PALETTE['gray2']}"; style="rounded";
        observ [label="Central Logs / Metrics / Traces"]; sec [label="Secrets / KMS / Certificate"]; cicd [label="CI/CD / IaC / Registry"]; dr [label="Cross-region Backup / DR"];
      }}
      internet -> waf -> lb; idp -> lb; lb -> webpod; lb -> realtime; webpod -> domain; domain -> worker;
      domain -> dbp; domain -> redis; domain -> mq; domain -> storage; domain -> search; worker -> dbp; worker -> mq; realtime -> mq;
      dbp -> dbr; dbp -> olap [label="CDC"]; mq -> olap; storage -> dr; dbp -> dr;
      observ -> domain [style=dashed]; sec -> domain [style=dashed]; cicd -> webpod [style=dashed]; cicd -> domain [style=dashed];
    }}'''
    out['04_deployment'] = render_dot('04_deployment', dot)

    dot = graph_header('TB', 0.38, 0.55, 'spline')
    dot += f'''
      source [label="业务命令\\nAPI / UI / 导入 / 设备", fillcolor="{PALETTE['sky']}"];
      aggregate [label="领域聚合\\n校验 · 状态机 · 事务", fillcolor="#DCE9FF", color="{PALETTE['blue']}"];
      db [label="业务表 + Outbox\\n同一事务提交", fillcolor="{PALETTE['green']}"];
      relay [label="Outbox Relay\\n发布并记录偏移"];
      bus [label="事件总线\\n按租户 / 领域 / 聚合键分区", fillcolor="{PALETTE['amber']}", color="#C58B00", width=3.5];
      inbox [label="消费者 Inbox\\n幂等去重", fillcolor="{PALETTE['green']}"];
      consumers [label="WMS / TMS / AMS / 结算 / 控制塔\\n异步订阅、Saga 与补偿"];
      read [label="查询投影\\n时间线 / 搜索 / KPI / 大屏", fillcolor="{PALETTE['violet']}"];
      dlq [label="重试 / 延迟队列 / 死信\\n人工修复与重放", fillcolor="{PALETTE['rose']}", color="{PALETTE['red']}"];
      source -> aggregate -> db -> relay -> bus;
      bus -> inbox; bus -> read; inbox -> consumers; inbox -> dlq [label="处理失败"];
      dlq -> inbox [label="修复后重放", style=dashed, constraint=false]; consumers -> bus [label="后续领域事件", style=dashed, constraint=false];
      {{rank=same; inbox; read;}}
    }}'''
    out['05_event_data'] = render_dot('05_event_data', dot)

    dot = graph_header('TB', 0.30, 0.45, 'spline')
    dot += f'''
      req [label="请求上下文\\nuserId · tenantId · orgId · warehouseId · locale", fillcolor="{PALETTE['sky']}"];
      authn [label="认证 AuthN\\nSSO / MFA / Token / API Credential", fillcolor="#DCE9FF"];
      authz [label="授权 AuthZ\\nRBAC + ABAC + 数据范围", fillcolor="#DCE9FF"];
      policy [label="策略判定\\n资源 · 动作 · 组织树 · 货主 · 仓库 · 数据所有者"];
      scope [label="租户隔离\\n行级 tenant_id / Schema / 独立库可选", fillcolor="{PALETTE['green']}"];
      mask [label="字段脱敏与导出控制\\n手机号 / 身份证 / 地址 / 价格"];
      audit [label="不可抵赖审计\\n登录 · 查询 · 变更 · 导出 · 审批 · 接口", fillcolor="{PALETTE['amber']}"];
      kms [label="密钥与加密\\nTLS · KMS · Secret Rotation · 对象存储加密"];
      req -> authn -> authz -> policy -> scope -> mask -> audit;
      kms -> authn [style=dashed]; kms -> scope [style=dashed]; kms -> audit [style=dashed];
    }}'''
    out['06_tenancy_security'] = render_dot('06_tenancy_security', dot)

    dot = graph_header('TB', 0.25, 0.42, 'spline')
    dot += f'''
      shell [label="Application Shell\\n左侧模块导航 · 顶部上下文 · 多标签工作区", fillcolor="{PALETTE['navy']}", fontcolor="white", color="{PALETTE['navy']}"];
      list [label="列表查询页\\nQueryPanel + CommandBar + DataGrid + SavedView", fillcolor="{PALETTE['sky']}"];
      detail [label="主从详情页\\n摘要 / 状态阶梯 / 子表 / 时间线 / 附件", fillcolor="{PALETTE['green']}"];
      task [label="任务工作台\\n队列 / 领取 / 扫描 / 异常 / 绩效", fillcolor="{PALETTE['amber']}"];
      board [label="地图与看板\\nKPI / 轨迹 / 热力 / 告警 / 下钻", fillcolor="{PALETTE['violet']}"];
      drawer [label="统一交互原语\\nDrawer / Modal / ContextMenu / Toast / Print / Export"];
      shell -> list; shell -> detail; shell -> task; shell -> board;
      list -> drawer; detail -> drawer; task -> drawer; board -> drawer;
    }}'''
    out['07_frontend_structure'] = render_dot('07_frontend_structure', dot)
    return out


def flow_dot(title: str, nodes: List[Tuple[str,str,str]], edges: List[Tuple[str,str,str]], rankdir='TB') -> str:
    dot = graph_header(rankdir, 0.30, 0.48, 'spline')
    dot += f'graph [label="{q(title)}", fontsize=15];\n'
    for key, label, kind in nodes:
        attrs = []
        if kind == 'start': attrs += ['shape=oval', f'fillcolor="{PALETTE["sky"]}"', f'color="{PALETTE["blue"]}"']
        elif kind == 'end': attrs += ['shape=oval', f'fillcolor="{PALETTE["green"]}"', f'color="{PALETTE["teal"]}"']
        elif kind == 'decision': attrs += ['shape=diamond', f'fillcolor="{PALETTE["amber"]}"', 'color="#C58B00"']
        elif kind == 'exception': attrs += [f'fillcolor="{PALETTE["rose"]}"', f'color="{PALETTE["red"]}"']
        elif kind == 'system': attrs += [f'fillcolor="{PALETTE["violet"]}"', f'color="{PALETTE["purple"]}"']
        else: attrs += [f'fillcolor="{PALETTE["white"]}"']
        dot += f'{key} [label="{q(label)}", {", ".join(attrs)}];\n'
    for a,b,label in edges:
        dot += f'{a} -> {b}' + (f' [label="{q(label)}"]' if label else '') + ';\n'
    dot += '}'
    return dot


def flow_diagrams() -> Dict[str, Path]:
    specs = {
      '10_e2e_order': ('端到端：订单到交付与结算',
        [('s','多渠道订单进入','start'),('v','OMS 校验、幂等与审核','process'),('a','库存/仓库/伙伴分配','decision'),('w','WMS 入库或出库履约','system'),('t','TMS 计划、委托与运输','system'),('p','签收 / POD / 差异','decision'),('b','计费、对账、财务过账','system'),('e','订单关闭与指标归档','end'),('x','例外工单与补偿','exception')],
        [('s','v',''),('v','a',''),('a','w','仓储履约'),('a','t','直运/运输需求'),('w','t','可提货/发运'),('t','p',''),('p','b','签收确认'),('b','e',''),('v','x','校验失败'),('w','x','短缺/差异'),('t','x','延误/拒单'),('x','v','修复/重试')]),
      '11_inbound': ('WMS 入库：预报到上架',
        [('s','采购单/ASN/调拨预报','start'),('ap','预约与到场签到','process'),('r','收货任务与扫描','process'),('d','数量/包装/批次是否一致','decision'),('q','质检/隔离/处置','decision'),('l','生成 LPN 与上架任务','process'),('p','按策略推荐库位','system'),('c','上架确认并增加可用库存','end'),('x','差异、破损、拒收、补收','exception')],
        [('s','ap',''),('ap','r',''),('r','d',''),('d','q','一致/需检'),('d','x','差异'),('q','l','放行'),('q','x','不合格'),('l','p',''),('p','c',''),('x','r','补收/复检')]),
      '12_outbound': ('WMS 出库：波次到发运',
        [('s','出库单释放','start'),('w','波次选择与订单分组','process'),('a','库存分配与缺货判断','decision'),('r','补货或重新分配','process'),('p','生成拣选任务/路径','process'),('v','复核、包装、称重、贴标','process'),('g','集货、装车与封签','process'),('h','发运确认/交接 TMS','end'),('x','短拣、破损、错拣','exception')],
        [('s','w',''),('w','a',''),('a','p','库存足够'),('a','r','不足'),('r','a','补货完成'),('p','v',''),('p','x','异常'),('x','p','复核后继续'),('v','g',''),('g','h','')]),
      '13_tms': ('TMS：计划、委托、执行与回单',
        [('s','运输需求/订单','start'),('pl','合单、设备与路线计划','process'),('ap','计划审批','decision'),('te','承运商委托/竞价/报价','process'),('ac','承运商是否接受','decision'),('d','车辆司机指派与发运','process'),('tr','GPS/APP/EDI 跟踪','system'),('pod','到达、签收与 POD 审核','process'),('e','结算事实与运输关闭','end'),('x','拒单、延误、偏航、货损','exception')],
        [('s','pl',''),('pl','ap',''),('ap','te','通过'),('ap','pl','驳回'),('te','ac',''),('ac','d','接受'),('ac','pl','拒绝/超时'),('d','tr',''),('tr','pod',''),('tr','x','异常'),('x','tr','处置后继续'),('pod','e','')]),
      '14_appointment': ('AMS：预约、排队、靠台与出场',
        [('s','关联订单或无单预约','start'),('cap','计算工作量与可用时隙','system'),('lock','并发锁定容量','decision'),('app','提交/自动确认/审批','process'),('rem','提醒与到场准备','process'),('in','门岗签到与证件校验','decision'),('q','排队叫号与月台分配','process'),('op','装卸作业事件','process'),('out','出场、关闭与绩效统计','end'),('x','冲突、迟到、爽约、超时','exception')],
        [('s','cap',''),('cap','lock',''),('lock','app','锁定成功'),('lock','x','容量冲突'),('app','rem','确认'),('rem','in',''),('in','q','校验通过'),('in','x','拒绝/异常'),('q','op','靠台'),('op','out',''),('rem','x','未到/迟到')]),
      '15_inventory_transfer': ('库存调整、移库与盘点',
        [('s','创建盘点/移库/调整请求','start'),('fr','冻结相关库存范围','process'),('scan','扫描实物与库位','process'),('diff','账实是否一致','decision'),('approve','差异审批与原因码','process'),('post','写库存流水与余额','system'),('rel','释放冻结/生成补货','end'),('x','并发占用或序列号冲突','exception')],
        [('s','fr',''),('fr','scan',''),('scan','diff',''),('diff','rel','一致'),('diff','approve','有差异'),('approve','post','通过'),('post','rel',''),('fr','x','无法冻结'),('x','s','修复后重建')]),
      '16_return': ('退货与逆向物流',
        [('s','客户退货/RMA申请','start'),('au','规则校验与退货授权','decision'),('tr','预约或逆向运输','process'),('rec','退货收货与识别','process'),('qc','质检与处置判定','decision'),('disp','上架/维修/报废/退供应商','process'),('fin','退款、换货与费用结算','process'),('e','RMA 关闭','end'),('x','无授权、超期、货不符','exception')],
        [('s','au',''),('au','tr','通过'),('au','x','拒绝'),('tr','rec',''),('rec','qc',''),('qc','disp',''),('qc','x','异常'),('disp','fin',''),('fin','e','')]),
      '17_crossdock': ('越库：入库直连出库',
        [('s','匹配入库供应与出库需求','start'),('m','按商品/批次/数量/时间窗匹配','system'),('ok','是否满足越库约束','decision'),('recv','收货到越库暂存区','process'),('alloc','绑定出库明细与运输班次','process'),('load','直接集货/装车','process'),('e','发运并完成双向状态','end'),('x','超时/短缺/质检不通过','exception')],
        [('s','m',''),('m','ok',''),('ok','recv','满足'),('ok','x','不满足'),('recv','alloc',''),('alloc','load',''),('load','e',''),('x','s','转普通入库/补货')]),
      '18_exception': ('跨域例外处理与补偿',
        [('s','业务事件或监控检测异常','start'),('r','规则匹配严重度与 SLA','system'),('case','创建例外工单并路由责任人','process'),('act','通知、冻结、重试或人工处理','decision'),('comp','执行补偿/改派/调整/重放','process'),('verify','验证业务状态与数据一致性','decision'),('e','关闭工单并沉淀根因','end'),('esc','升级、加急和管理层告警','exception')],
        [('s','r',''),('r','case',''),('case','act',''),('act','comp','可处置'),('act','esc','超时/高严重度'),('esc','comp',''),('comp','verify',''),('verify','e','通过'),('verify','case','未修复')]),
      '19_settlement': ('计费、对账、开票与收付款',
        [('s','业务完成产生计费事实','start'),('rate','匹配合同/费率/税率版本','system'),('calc','计算阶梯、最低费、附加费','process'),('val','规则与金额校验','decision'),('v','生成应收/应付/预提凭证','process'),('rec','对账与差异协同','decision'),('inv','开票/收票与财务过账','process'),('pay','收付款、核销与期间关闭','end'),('x','费率缺失、争议、调整','exception')],
        [('s','rate',''),('rate','calc',''),('calc','val',''),('val','v','通过'),('val','x','失败'),('v','rec',''),('rec','inv','确认'),('rec','x','争议'),('x','calc','调整后重算'),('inv','pay','')])
    }
    out = {}
    for key,(title,nodes,edges) in specs.items():
        out[key] = render_dot(key, flow_dot(title,nodes,edges))
    return out


def draw_sequence(name: str, actors: List[str], messages: List[Tuple[int,int,str,str]], title: str) -> Path:
    regular = font_manager.FontProperties(fname=FONT_PATH)
    bold = font_manager.FontProperties(fname=FONT_BOLD_PATH)
    n, steps = len(actors), len(messages)
    fig, ax = plt.subplots(figsize=(max(11, n*1.75), max(6, steps*0.42+2.0)))
    ax.set_xlim(-0.6,n-0.4); ax.set_ylim(steps+1.8,-1); ax.axis('off')
    for i,a in enumerate(actors):
        rect=plt.Rectangle((i-.43,-.55),.86,.62,linewidth=1.2,edgecolor=PALETTE['blue'],facecolor=PALETTE['sky'])
        ax.add_patch(rect); ax.text(i,-.24,a,ha='center',va='center',fontsize=9.7,fontproperties=bold)
        ax.plot([i,i],[.08,steps+1.05],linestyle='--',linewidth=.8,color='#9CA3AF')
    y=.75
    for idx,(src,dst,label,style) in enumerate(messages,1):
        if style=='note':
            ax.text((src+dst)/2,y,label,ha='center',va='center',fontsize=8.2,fontproperties=regular,
                    bbox=dict(boxstyle='round,pad=.25',facecolor=PALETTE['amber'],edgecolor='#C58B00'))
            y+=.75; continue
        ls='--' if style=='return' else '-'
        ax.annotate('',xy=(dst,y),xytext=(src,y),arrowprops=dict(arrowstyle='-|>',lw=1.1,linestyle=ls,color=PALETTE['navy'],shrinkA=5,shrinkB=5))
        mid=(src+dst)/2
        ax.text(mid,y-.13,f'{idx}. {label}',ha='center',va='bottom',fontsize=8.2,fontproperties=regular,
                bbox=dict(boxstyle='round,pad=.10',facecolor='white',edgecolor='none',alpha=.94))
        if style=='async': ax.text(mid,y+.18,'异步',ha='center',va='top',fontsize=6.8,color=PALETTE['orange'],fontproperties=regular)
        y+=.75
    ax.set_title(title,fontsize=14,pad=12,fontproperties=bold)
    plt.tight_layout()
    png=DIAG/f'{name}.png'; svg=DIAG/f'{name}.svg'
    fig.savefig(png,dpi=220,bbox_inches='tight',facecolor='white'); fig.savefig(svg,bbox_inches='tight',facecolor='white'); plt.close(fig)
    return png


def sequence_diagrams() -> Dict[str,Path]:
    return {
      '20_order_api_seq': draw_sequence('20_order_api_seq',['ERP/电商','API网关','OMS','事件总线','WMS','TMS','控制塔'],[
        (0,1,'POST /orders + Idempotency-Key',''),(1,2,'鉴权、租户解析、创建订单',''),(2,2,'校验、去重、状态机与 Outbox','note'),(2,0,'202 Accepted / orderId','return'),(2,3,'order.released.v1','async'),(3,4,'创建出库/入库履约任务','async'),(3,5,'创建运输需求','async'),(3,6,'投影订单时间线与 SLA','async'),(4,3,'outbound.ready.v1','async'),(3,5,'更新可提货状态','async'),(5,3,'shipment.delivered.v1','async'),(3,2,'回写履约与完成状态','async')], '时序图：外部订单接入与跨域异步编排'),
      '21_wms_rf_seq': draw_sequence('21_wms_rf_seq',['RF设备','移动BFF','WMS任务','库存聚合','打印服务','事件总线'],[
        (0,1,'领取拣选任务',''),(1,2,'校验用户/仓库/设备',''),(2,0,'任务与推荐路径','return'),(0,1,'扫描库位、商品、LPN、数量',''),(1,2,'提交扫描命令 + taskVersion',''),(2,3,'锁定/扣减/移动库存',''),(3,2,'库存新版本','return'),(2,2,'记录任务步骤与离线序号','note'),(2,4,'需要时打印箱标/补标','async'),(2,5,'pick.confirmed.v1','async'),(2,0,'下一步/完成/异常提示','return')], '时序图：RF 扫描、库存并发控制与任务确认'),
      '22_driver_seq': draw_sequence('22_driver_seq',['司机APP','移动BFF','TMS','地图/GPS','事件总线','客户门户','控制塔'],[
        (0,1,'接受任务并检查车辆证照',''),(1,2,'确认指派',''),(0,1,'发运/到达/装卸/签收节点',''),(1,2,'校验节点顺序与时间窗',''),(2,3,'获取位置、路线与 ETA',''),(3,2,'位置/ETA','return'),(2,4,'tracking.event.v1','async'),(4,5,'更新客户可视状态','async'),(4,6,'更新大屏与 SLA','async'),(0,1,'上传签名、照片、回单',''),(1,2,'POD 元数据 + 对象引用',''),(2,4,'shipment.delivered.v1','async')], '时序图：司机执行、轨迹事件、客户可视化与 POD'),
      '23_appointment_seq': draw_sequence('23_appointment_seq',['客户/承运商','AMS','容量服务','审批流','门岗','月台作业','事件总线'],[
        (0,1,'查询可约日期与工作量',''),(1,2,'计算时隙容量',''),(2,1,'候选时隙 + 版本','return'),(0,1,'提交预约 slotId + version',''),(1,2,'原子占用容量',''),(2,1,'占用成功/冲突','return'),(1,3,'必要时发起审批',''),(3,1,'批准/改期/拒绝','return'),(1,6,'appointment.confirmed.v1','async'),(0,4,'二维码/证件签到',''),(4,1,'check-in 事件',''),(1,5,'排队叫号与月台分配',''),(5,1,'开始/暂停/完成装卸',''),(1,6,'appointment.completed.v1','async')], '时序图：预约容量并发、审批、门岗和月台执行'),
      '24_billing_seq': draw_sequence('24_billing_seq',['业务域','事件总线','计费服务','合同/费率','结算服务','客户/承运商','财务/ERP'],[
        (0,1,'业务完成事件（含计费事实）','async'),(1,2,'订阅并幂等落地',''),(2,3,'按生效时间匹配合同与费率版本',''),(3,2,'规则、阶梯、最低费、税率','return'),(2,2,'生成可追溯计算明细','note'),(2,4,'创建应收/应付/预提凭证',''),(4,5,'发布对账单/差异协同',''),(5,4,'确认或提出争议',''),(4,4,'调整、分摊并重新审批','note'),(4,6,'过账、开票、收付款',''),(6,4,'回写财务状态','return')], '时序图：事件驱动的计费、对账与财务过账')
    }


def state_diagrams() -> Dict[str,Path]:
    specs = {
      '30_order_state': ('订单状态机',[('Draft','草稿'),('Open','已创建'),('Approved','已审核'),('Allocated','已分配'),('Released','已释放'),('Executing','履约中'),('Completed','已完成'),('Closed','已关闭'),('Hold','已冻结'),('Cancelled','已取消')],[('Draft','Open','提交'),('Open','Approved','审核'),('Approved','Allocated','分配'),('Allocated','Released','释放'),('Released','Executing','首个执行事件'),('Executing','Completed','全部交付'),('Completed','Closed','结算/归档'),('Open','Hold','冻结'),('Approved','Hold','冻结'),('Hold','Open','解冻/回退'),('Open','Cancelled','取消'),('Approved','Cancelled','撤销后取消')]),
      '31_inbound_state': ('WMS 入库状态机',[('Draft','草稿'),('Expected','待到货'),('Arrived','已到场'),('Receiving','收货中'),('Received','已收货'),('Inspecting','质检中'),('Putaway','上架中'),('Completed','已完成'),('Exception','异常'),('Cancelled','已取消')],[('Draft','Expected','发布'),('Expected','Arrived','签到'),('Arrived','Receiving','开始收货'),('Receiving','Received','收货完成'),('Received','Inspecting','需质检'),('Received','Putaway','免检'),('Inspecting','Putaway','放行'),('Putaway','Completed','全部上架'),('Receiving','Exception','差异'),('Inspecting','Exception','不合格'),('Exception','Receiving','补收'),('Exception','Putaway','处置放行'),('Draft','Cancelled','取消')]),
      '32_outbound_state': ('WMS 出库状态机',[('Draft','草稿'),('Released','已释放'),('Waved','已入波次'),('Allocated','已分配'),('Picking','拣选中'),('Picked','已拣完'),('Packed','已包装'),('Staged','已集货'),('Loaded','已装车'),('Shipped','已发运'),('Shortage','缺货异常'),('Cancelled','已取消')],[('Draft','Released','发布'),('Released','Waved','入波次'),('Waved','Allocated','分配'),('Allocated','Picking','领取任务'),('Picking','Picked','拣完'),('Picked','Packed','复核包装'),('Packed','Staged','集货'),('Staged','Loaded','装车'),('Loaded','Shipped','发运'),('Waved','Shortage','不足'),('Picking','Shortage','短拣'),('Shortage','Allocated','补货后重分配'),('Draft','Cancelled','取消')]),
      '33_shipment_state': ('TMS 运输单状态机',[('Open','已创建'),('Planned','已计划'),('Approved','已审批'),('Tendered','已委托'),('Accepted','已接单'),('Dispatched','已派车'),('Tracking','运输中'),('Arrived','已到达'),('Delivered','已签收'),('POD','回单确认'),('Settled','已结算'),('Closed','已关闭'),('Rejected','已拒绝'),('Cancelled','已取消')],[('Open','Planned','计划'),('Planned','Approved','审批'),('Approved','Tendered','委托'),('Tendered','Accepted','接受'),('Tendered','Rejected','拒绝/超时'),('Rejected','Planned','改派'),('Accepted','Dispatched','派车'),('Dispatched','Tracking','发运'),('Tracking','Arrived','到达'),('Arrived','Delivered','签收'),('Delivered','POD','回单审核'),('POD','Settled','结算'),('Settled','Closed','归档'),('Open','Cancelled','取消')]),
      '34_appointment_state': ('预约状态机',[('Draft','草稿'),('Submitted','已提交'),('Pending','待审批'),('Confirmed','已确认'),('Rescheduled','已改期'),('CheckedIn','已签到'),('Queued','排队中'),('Docked','已靠台'),('Operating','作业中'),('CheckedOut','已出场'),('Completed','已完成'),('NoShow','爽约'),('Rejected','已拒绝'),('Cancelled','已取消')],[('Draft','Submitted','提交'),('Submitted','Pending','需审批'),('Submitted','Confirmed','自动确认'),('Pending','Confirmed','批准'),('Pending','Rejected','拒绝'),('Confirmed','Rescheduled','改期'),('Rescheduled','Confirmed','确认新时段'),('Confirmed','CheckedIn','签到'),('CheckedIn','Queued','排队'),('CheckedIn','Docked','直接靠台'),('Queued','Docked','叫号'),('Docked','Operating','开始装卸'),('Operating','CheckedOut','完成出场'),('CheckedOut','Completed','关闭'),('Confirmed','NoShow','超时未到'),('Draft','Cancelled','取消')]),
      '35_voucher_state': ('结算凭证状态机',[('Draft','草稿'),('Calculated','已计算'),('Validated','已校验'),('Approved','已审核'),('Reconciled','已对账'),('Invoiced','已开/收票'),('Paid','已收/付款'),('Closed','已关闭'),('Disputed','争议中'),('Adjusted','已调整'),('Voided','已作废')],[('Draft','Calculated','计费'),('Calculated','Validated','校验'),('Validated','Approved','审核'),('Approved','Reconciled','对账确认'),('Approved','Disputed','争议'),('Disputed','Adjusted','调整'),('Adjusted','Validated','重校验'),('Reconciled','Invoiced','开/收票'),('Invoiced','Paid','收付款'),('Paid','Closed','期间关闭'),('Draft','Voided','作废')])
    }
    out={}
    for key,(title,nodes,edges) in specs.items():
        dot=graph_header('LR',0.22,0.48,'spline')+f'graph [label="{q(title)}", fontsize=15];\n'
        for k,l in nodes:
            fill=PALETTE['sky'] if k in ('Draft','Open','Expected') else (PALETTE['green'] if k in ('Closed','Completed','Shipped','Settled') else (PALETTE['rose'] if k in ('Exception','Shortage','Rejected','Cancelled','NoShow','Voided') else 'white'))
            dot+=f'{k} [label="{q(l)}", fillcolor="{fill}"];\n'
        for a,b,l in edges: dot+=f'{a} -> {b} [label="{q(l)}"];\n'
        dot+='}'
        out[key]=render_dot(key,dot)
    return out


def class_node(key: str, title: str, fields: List[str], fill='white') -> str:
    rows=''.join(f'<TR><TD ALIGN="LEFT"><FONT POINT-SIZE="9">{q(f)}</FONT></TD></TR>' for f in fields)
    return f'''{key} [shape=plain, label=<<TABLE BORDER="1" CELLBORDER="0" CELLSPACING="0" CELLPADDING="5" COLOR="#AAB4C3">
      <TR><TD BGCOLOR="{fill}"><B>{q(title)}</B></TD></TR>{rows}</TABLE>>];\n'''


def class_diagrams() -> Dict[str,Path]:
    out={}
    # Context overview
    dot=graph_header('TB',0.30,0.55,'spline')+'graph [label="领域聚合与上下文边界", fontsize=15];\n'
    contexts=[('IAM','IAM / Tenant / Org'),('MDM','Master Data / Config'),('OMS','OMS 订单'),('WMS','WMS 仓储'),('TMS','TMS 运输'),('AMS','AMS 预约'),('BILL','Billing / Settlement'),('CTL','Control Tower / BI'),('INT','Integration / Document')]
    for k,l in contexts:
        fill=PALETTE['sky'] if k in ('OMS','WMS','TMS','AMS') else (PALETTE['green'] if k in ('IAM','MDM') else PALETTE['violet'])
        dot+=f'{k} [label="{l}", fillcolor="{fill}", width=2.4, height=.65];\n'
    for a,b,l in [('IAM','OMS','授权上下文'),('MDM','OMS','主数据快照'),('OMS','WMS','履约单/事件'),('OMS','TMS','运输需求'),('OMS','AMS','预约需求'),('WMS','TMS','可提货/装车'),('AMS','WMS','到场/靠台'),('WMS','BILL','仓储计费事实'),('TMS','BILL','运输计费事实'),('AMS','BILL','预约/爽约事实'),('OMS','CTL','订单事件'),('WMS','CTL','库存与作业事件'),('TMS','CTL','轨迹与异常'),('INT','OMS','API/EDI'),('INT','TMS','GPS/POD')]:
        dot+=f'{a} -> {b} [label="{l}"];\n'
    dot+='}'; out['40_class_overview']=render_dot('40_class_overview',dot)

    dot=graph_header('TB',0.18,0.42,'spline')+'graph [label="类图：租户、权限与主数据", fontsize=15];\n'
    for args in [
      ('Tenant','Tenant 租户',['id: UUID','code/name: String','status: TenantStatus','timezone/currency: String','configVersion: Long'],PALETTE['sky']),
      ('Organization','Organization 组织',['id/tenantId/parentId: UUID','type: OrgType','path: String','status: Status'],'white'),
      ('User','User 用户',['id/tenantId: UUID','loginName/displayName: String','status: UserStatus','identityProvider: String'],'white'),
      ('Role','Role 角色',['id/tenantId: UUID','code/name: String','status: Status'],'white'),
      ('Permission','Permission 权限',['id: UUID','resource/action/field: String','effect: ALLOW/DENY'],'white'),
      ('DataScope','DataScope 数据范围',['id: UUID','roleId: UUID','dimension: ORG/WAREHOUSE/OWNER','expression: JSON'],'white'),
      ('Product','Product 商品',['id/tenantId: UUID','sku/name: String','baseUom: String','lot/serial rules: JSON','status: Status'],PALETTE['green']),
      ('Partner','Partner 伙伴',['id/tenantId: UUID','type: CUSTOMER/SUPPLIER/CARRIER','code/name: String','credit/certificates: JSON'],PALETTE['green']),
      ('Warehouse','Warehouse 仓库',['id/orgId: UUID','code/name: String','addressId: UUID','timezone: String'],PALETTE['green']),
      ('Location','Location 库位/月台',['id/warehouseId: UUID','code/type/zone: String','capacity: JSON','status: Status'],PALETTE['green']),
      ('Contract','Contract 合同',['id/partyId: UUID','type: SALES/PURCHASE/FREIGHT/WAREHOUSE','validFrom/To: Date','status: Status'],PALETTE['green']),
      ('RateCard','RateCard 费率',['id/contractId: UUID','version: String','currency/taxMode: String','effectiveAt: Instant'],PALETTE['green'])]:
        dot+=class_node(*args)
    for a,b,l in [('Tenant','Organization','1..*'),('Tenant','User','1..*'),('User','Role','*..*'),('Role','Permission','*..*'),('Role','DataScope','1..*'),('Organization','Warehouse','1..*'),('Warehouse','Location','1..*'),('Partner','Contract','0..*'),('Contract','RateCard','1..*')]: dot+=f'{a} -> {b} [label="{l}", arrowhead=none];\n'
    dot+='}'; out['41_class_iam_master']=render_dot('41_class_iam_master',dot)

    dot=graph_header('TB',0.16,0.42,'spline')+'graph [label="类图：OMS 订单与履约", fontsize=15];\n'
    classes=[
      ('BusinessOrder','BusinessOrder 业务订单',['id/tenantId: UUID','orderNo/externalNo: String','type/channel: Enum','customerId/supplierId: UUID?','status: OrderStatus','priority: Int','requestedAt: Instant','version: Long'],PALETTE['sky']),
      ('OrderLine','OrderLine 订单行',['id/orderId: UUID','lineNo: Int','productId: UUID','orderedQty: Decimal','unitPrice: Money?','lot/serial requirements: JSON','status: LineStatus'],'white'),
      ('OrderAddress','OrderAddress',['id/orderId: UUID','type: SHIP_FROM/TO/BILL','addressId: UUID','contactSnapshot: JSON'],'white'),
      ('OrderParty','OrderParty',['id/orderId: UUID','role: PartnerRole','partnerId: UUID','snapshot: JSON'],'white'),
      ('OrderHold','OrderHold 冻结',['id/orderId: UUID','reasonCode: String','holdAt/releasedAt: Instant','createdBy: UUID'],'white'),
      ('OrderEvent','OrderEvent 时间线',['id/orderId: UUID','eventType: String','occurredAt: Instant','source/payload: JSON'],'white'),
      ('Allocation','Allocation 分配',['id/orderLineId: UUID','warehouseId/ownerId: UUID','allocatedQty: Decimal','ruleId: UUID?','status: AllocationStatus'],'white'),
      ('Fulfillment','FulfillmentOrder 履约单',['id/fulfillmentNo: UUID/String','sourceOrderId/warehouseId: UUID','type: INBOUND/OUTBOUND/TRANSFER','status: FulfillmentStatus','plannedDate: Date'],'white'),
      ('FulfillmentLine','FulfillmentLine',['id/fulfillmentId/sourceLineId: UUID','productId: UUID','qty: Decimal','status: Status'],'white'),
      ('ShipmentRequest','ShipmentRequest 运输需求',['id/requestNo: UUID/String','orderId: UUID','origin/destination: UUID','pickup/deliveryWindow: TimeWindow','weight/volume: Decimal','status: RequestStatus'],'white'),
      ('RMA','ReturnAuthorization 退货授权',['id/rmaNo: UUID/String','originalOrderId: UUID','reasonCode: String','status: RmaStatus','resolution: Refund/Replace/Repair'],'white'),
      ('SettlementRequest','SettlementRequest 结算请求',['id/orderId: UUID','chargeFacts: JSON','status: Status'],'white')]
    for c in classes: dot+=class_node(*c)
    for a,b,l in [('BusinessOrder','OrderLine','1..*'),('BusinessOrder','OrderAddress','1..*'),('BusinessOrder','OrderParty','1..*'),('BusinessOrder','OrderHold','0..*'),('BusinessOrder','OrderEvent','0..*'),('OrderLine','Allocation','0..*'),('BusinessOrder','Fulfillment','0..*'),('Fulfillment','FulfillmentLine','1..*'),('BusinessOrder','ShipmentRequest','0..*'),('BusinessOrder','RMA','0..*'),('BusinessOrder','SettlementRequest','0..*')]: dot+=f'{a} -> {b} [label="{l}", arrowhead=none];\n'
    dot+='}'; out['42_class_oms']=render_dot('42_class_oms',dot)

    dot=graph_header('LR',0.16,0.40,'spline')+'graph [label="类图：WMS 入库与库存", fontsize=15];\n'
    classes=[
      ('Inbound','InboundOrder 入库单',['id/inboundNo: UUID/String','sourceOrderId/warehouseId/ownerId: UUID','type: PURCHASE/RETURN/TRANSFER','expectedAt: Instant','status: InboundStatus'],PALETTE['green']),
      ('InboundLine','InboundLine',['id/inboundId: UUID','productId: UUID','expected/received/acceptedQty: Decimal','lotRule: JSON','status: Status'],'white'),
      ('ReceiptTask','ReceiptTask 收货任务',['id/taskNo: UUID/String','inboundId/dockId/assigneeId: UUID?','started/completedAt: Instant','status: TaskStatus'],'white'),
      ('ReceiptLine','ReceiptLine',['id/receiptTaskId/inboundLineId: UUID','handlingUnitId/lotId: UUID?','receivedQty: Decimal','condition: String'],'white'),
      ('Inspection','Inspection 质检',['id/sourceLineId: UUID','planId: UUID','sampleQty/accepted/rejected: Decimal','result/status: Enum'],'white'),
      ('Putaway','PutawayTask 上架任务',['id/taskNo: UUID/String','handlingUnitId: UUID','from/toLocationId: UUID','qty: Decimal','strategyTrace: JSON','status: TaskStatus'],'white'),
      ('HandlingUnit','HandlingUnit LPN',['id/warehouseId: UUID','lpn: String','parentId: UUID?','type/status: Enum','weight/volume: Decimal'],'white'),
      ('Lot','InventoryLot 批次',['id/productId/ownerId: UUID','lotNo: String','mfg/expiryDate: Date','attributes: JSON'],'white'),
      ('Serial','SerialNumber 序列号',['id/productId: UUID','serialNo: String','status: SerialStatus','currentHu/locationId: UUID'],'white'),
      ('Balance','InventoryBalance 库存余额',['id/warehouse/location/product/owner/lot: UUID','onHand/available/allocated/hold: Decimal','version: Long'],PALETTE['sky']),
      ('Movement','InventoryMovement 库存流水',['id: UUID','type: RECEIVE/PUTAWAY/PICK/ADJUST/TRANSFER','businessRef: String','from/to: JSON','qty: Decimal','occurredAt: Instant'],'white'),
      ('Count','CountOrder 盘点单',['id/countNo: UUID/String','scope: JSON','freezePolicy: Enum','status: CountStatus'],'white'),
      ('CountLine','CountLine',['id/countId/balanceKey: UUID/JSON','bookQty/countQty/diffQty: Decimal','reasonCode: String'],'white'),
      ('Adjustment','Adjustment 库存调整',['id/adjustNo: UUID/String','reasonCode: String','lines: JSON','status: ApprovalStatus'],'white')]
    for c in classes: dot+=class_node(*c)
    for a,b,l in [('Inbound','InboundLine','1..*'),('Inbound','ReceiptTask','0..*'),('ReceiptTask','ReceiptLine','1..*'),('ReceiptLine','Inspection','0..1'),('ReceiptLine','HandlingUnit','0..1'),('HandlingUnit','Putaway','0..*'),('HandlingUnit','Balance','0..*'),('Lot','Balance','0..*'),('Serial','HandlingUnit','0..*'),('Balance','Movement','0..*'),('Count','CountLine','1..*'),('CountLine','Adjustment','0..1'),('Adjustment','Movement','1..*')]: dot+=f'{a} -> {b} [label="{l}", arrowhead=none];\n'
    dot+='}'; out['43_class_wms_inbound_inventory']=render_dot('43_class_wms_inbound_inventory',dot)

    dot=graph_header('TB',0.16,0.40,'spline')+'graph [label="类图：WMS 出库与作业", fontsize=15];\n'
    classes=[
      ('Outbound','OutboundOrder 出库单',['id/outboundNo: UUID/String','sourceOrderId/warehouseId/ownerId: UUID','type: SALES/TRANSFER/RETURN','shipWindow: TimeWindow','status: OutboundStatus'],PALETTE['green']),
      ('OutboundLine','OutboundLine',['id/outboundId/productId: UUID','requested/allocated/picked/shippedQty: Decimal','status: Status'],'white'),
      ('Wave','Wave 波次',['id/waveNo: UUID/String','templateId/warehouseId: UUID','plannedStartAt: Instant','status: WaveStatus'],'white'),
      ('WaveOrder','WaveOrder',['waveId/outboundId: UUID','sequence: Int'],'white'),
      ('AllocationDetail','AllocationDetail 分配明细',['id/outboundLineId/balanceId: UUID','allocatedQty: Decimal','strategyTrace: JSON'],'white'),
      ('PickTask','PickTask 拣选任务',['id/taskNo: UUID/String','waveId/assigneeId/deviceId: UUID?','routeSeq: Int','status: TaskStatus'],'white'),
      ('PickLine','PickTaskLine',['id/pickTaskId/allocationId: UUID','fromLocationId/targetHuId: UUID','qty/confirmedQty: Decimal'],'white'),
      ('PackTask','PackTask 包装任务',['id/packNo: UUID/String','outboundId/stationId: UUID','status: TaskStatus','weight/volume: Decimal'],'white'),
      ('Package','PackageUnit 包裹/LPN',['id/packageNo: UUID/String','outboundId/parentId: UUID?','labelNo: String','weight/volume: Decimal','status: PackageStatus'],'white'),
      ('Staging','StagingTask 集货任务',['id/outboundId/stagingLocationId: UUID','route/shipmentId: UUID?','status: TaskStatus'],'white'),
      ('Load','LoadTask 装车任务',['id/loadNo: UUID/String','shipmentId/dockId/vehicleId: UUID?','sealNo: String','status: TaskStatus'],'white'),
      ('VAS','ValueAddedOrder 增值服务',['id/vasNo: UUID/String','type: LABEL/REPACK/KIT/ASSEMBLY','sourceOrderId: UUID','status: VasStatus'],'white'),
      ('Labor','LaborTask 劳务任务',['id/taskType/businessTaskId: UUID/String','worker/teamId: UUID','standard/actualMinutes: Decimal','score: Decimal'],'white'),
      ('Replen','ReplenishmentTask 补货',['id/from/toLocationId: UUID','productId/lotId: UUID','qty: Decimal','triggerType/status: Enum'],'white')]
    for c in classes: dot+=class_node(*c)
    for a,b,l in [('Outbound','OutboundLine','1..*'),('Wave','WaveOrder','1..*'),('Outbound','WaveOrder','0..*'),('OutboundLine','AllocationDetail','0..*'),('Wave','PickTask','0..*'),('PickTask','PickLine','1..*'),('AllocationDetail','PickLine','1'),('Outbound','PackTask','0..*'),('PackTask','Package','1..*'),('Outbound','Staging','0..*'),('Staging','Load','0..*'),('Outbound','VAS','0..*'),('PickTask','Labor','0..*'),('AllocationDetail','Replen','0..1')]: dot+=f'{a} -> {b} [label="{l}", arrowhead=none];\n'
    dot+='}'; out['44_class_wms_outbound']=render_dot('44_class_wms_outbound',dot)

    dot=graph_header('TB',0.16,0.40,'spline')+'graph [label="类图：TMS 计划、执行与回单", fontsize=15];\n'
    classes=[
      ('TO','TransportOrder 运输订单',['id/transportNo: UUID/String','sourceOrderId: UUID?','mode: TransportMode','origin/destination: UUID','pickup/deliveryWindow: TimeWindow','status: TransportOrderStatus'],PALETTE['amber']),
      ('TOL','TransportOrderLine',['id/transportOrderId: UUID','productId: UUID?','packageCount/weight/volume: Decimal','temperatureBand: String'],'white'),
      ('Leg','TransportLeg 运输段',['id/transportOrderId: UUID','seq: Int','from/toStop: UUID','mode: TransportMode','status: LegStatus'],'white'),
      ('Plan','LoadPlan 配载计划',['id/planNo/planningBatchId: UUID/String','equipmentTypeId: UUID','weight/volumeUtilization: Decimal','status: PlanStatus'],'white'),
      ('Shipment','Shipment/Waybill 运单',['id/shipmentNo: UUID/String','loadPlanId/carrierId: UUID?','status: ShipmentStatus','planned/actual times: JSON'],'white'),
      ('ShipmentItem','ShipmentItem',['id/shipmentId/transportOrderLineId: UUID','qty/weight/volume: Decimal'],'white'),
      ('Stop','Stop 节点',['id/shipmentId: UUID','seq: Int','type: PICKUP/DELIVERY/HUB','addressId: UUID','planned/actual times: JSON','status: StopStatus'],'white'),
      ('Tender','CarrierTender 委托/投标',['id/shipmentId/carrierId: UUID','method: DIRECT/BID/QUOTE','offeredPrice: Money?','status: TenderStatus','expiresAt: Instant'],'white'),
      ('Assign','VehicleAssignment',['id/shipmentId/vehicleId/driverId: UUID','assignedAt: Instant','status: AssignmentStatus'],'white'),
      ('Route','RoutePlan 路线计划',['id/shipmentId: UUID','polyline: GeoJson','distance/duration: Decimal','optimizationTrace: JSON'],'white'),
      ('Track','TrackingEvent 跟踪事件',['id/shipmentId/stopId: UUID?','eventType: TrackingType','occurredAt: Instant','location: GeoPoint','source/payload: JSON'],'white'),
      ('Ex','TransportException 异常',['id/shipmentId: UUID','type/severity/status: String','ownerId: UUID?','detected/resolvedAt: Instant','resolution: JSON'],'white'),
      ('POD','ProofOfDelivery 回单',['id/shipmentId: UUID','signedAt/signer: String','signatureObjectId: UUID?','attachments: Set','status: PodStatus'],'white'),
      ('Fact','FreightChargeFact 计费事实',['id/shipmentId: UUID','distance/weight/volume/duration: Decimal','accessorials: JSON','confirmedAt: Instant'],'white')]
    for c in classes: dot+=class_node(*c)
    for a,b,l in [('TO','TOL','1..*'),('TO','Leg','1..*'),('Plan','Shipment','1..*'),('Shipment','ShipmentItem','1..*'),('TOL','ShipmentItem','0..*'),('Shipment','Stop','2..*'),('Shipment','Tender','0..*'),('Shipment','Assign','0..1'),('Shipment','Route','0..1'),('Shipment','Track','0..*'),('Shipment','Ex','0..*'),('Shipment','POD','0..1'),('Shipment','Fact','0..*')]: dot+=f'{a} -> {b} [label="{l}", arrowhead=none];\n'
    dot+='}'; out['45_class_tms']=render_dot('45_class_tms',dot)

    dot=graph_header('TB',0.18,0.42,'spline')+'graph [label="类图：AMS 预约、排队与月台", fontsize=15];\n'
    classes=[
      ('Calendar','CapacityCalendar 容量日历',['id/warehouseId/resourceId: UUID','resourceType: DOCK/ZONE/TEAM','date: Date','capacityProfileId: UUID'],PALETTE['cyan']),
      ('Slot','TimeSlot 时间窗',['id/calendarId: UUID','startAt/endAt: Instant','capacity/used/reserved: Decimal','status: SlotStatus','version: Long'],'white'),
      ('Appointment','Appointment 预约',['id/appointmentNo: UUID/String','type: INBOUND/OUTBOUND','warehouseId/slotId: UUID','carrier/driver/vehicleId: UUID?','workload: Decimal','status: AppointmentStatus'],'white'),
      ('OrderLink','AppointmentOrderLink',['appointmentId: UUID','orderType/orderId: String','plannedQty: Decimal'],'white'),
      ('Approval','ApprovalTask 审批',['id/appointmentId/assigneeId: UUID','decision: APPROVE/REJECT/RESCHEDULE','comment: String','decidedAt: Instant'],'white'),
      ('Queue','QueueTicket 排队号',['id/appointmentId: UUID','ticketNo: String','priority: Int','issuedAt/calledAt: Instant','status: QueueStatus'],'white'),
      ('Gate','GateEvent 门岗事件',['id/appointmentId: UUID','type: CHECK_IN/CHECK_OUT','occurredAt: Instant','operator/deviceId: UUID','verification: JSON'],'white'),
      ('Pass','GatePass 通行证',['id/appointmentId: UUID','qrCode: String','validFrom/To: Instant','status: PassStatus'],'white'),
      ('Dock','DockAssignment 月台分配',['id/appointmentId/dockId: UUID','assignedAt/releasedAt: Instant','status: AssignmentStatus'],'white'),
      ('Op','OperationEvent 装卸事件',['id/appointmentId/dockAssignmentId: UUID','type: START/PAUSE/COMPLETE','occurredAt: Instant','qty/payload: JSON'],'white'),
      ('NoShow','NoShowCase 爽约',['id/appointmentId: UUID','detectedAt: Instant','reasonCode: String','penaltyFactId: UUID?','status: Status'],'white')]
    for c in classes: dot+=class_node(*c)
    for a,b,l in [('Calendar','Slot','1..*'),('Slot','Appointment','0..*'),('Appointment','OrderLink','0..*'),('Appointment','Approval','0..*'),('Appointment','Queue','0..1'),('Appointment','Pass','0..1'),('Appointment','Gate','0..*'),('Appointment','Dock','0..1'),('Dock','Op','0..*'),('Appointment','NoShow','0..1')]: dot+=f'{a} -> {b} [label="{l}", arrowhead=none];\n'
    dot+='}'; out['46_class_ams']=render_dot('46_class_ams',dot)

    dot=graph_header('LR',0.14,0.36,'spline')+'graph [label="类图：计费、控制塔、工作流与集成", fontsize=15];\n'
    classes=[
      ('Rule','ChargeRule 计费规则',['id/rateCardId: UUID','chargeCode/scope: String','basis: QTY/WEIGHT/VOLUME/DISTANCE/TIME','expression: DSL/JSON','priority/status: Int/Enum'], '#F7EAF8'),
      ('Calc','ChargeCalculation 计算',['id/businessType/businessId: UUID/String','ruleId/rateVersion: UUID','inputFacts: JSON','amount/tax: Money','calculationTrace: JSON','status: CalcStatus'],'white'),
      ('Voucher','SettlementVoucher 凭证',['id/voucherNo: UUID/String','type: AR/AP/ACCRUAL/ADJUSTMENT','partyId/period: UUID/String','amount/taxAmount: Money','status: VoucherStatus'],'white'),
      ('VLine','VoucherLine',['id/voucherId/calcId: UUID','chargeCode/businessRef: String','amount/tax: Money'],'white'),
      ('Recon','ReconciliationStatement 对账单',['id/statementNo: UUID/String','partyId/period: UUID/String','totalAmount: Money','status: ReconciliationStatus'],'white'),
      ('Invoice','Invoice 发票',['id/invoiceNo: UUID/String','statementId: UUID','direction: AR/AP','amount/tax: Money','status: InvoiceStatus'],'white'),
      ('Payment','Payment 收付款',['id/invoiceId: UUID','paymentNo: String','amount: Money','paidAt: Instant','status: PaymentStatus'],'white'),
      ('Event','BusinessEvent 业务事件',['eventId/tenantId: UUID','aggregateType/id/version: String','eventType/traceId: String','occurredAt: Instant','payload: JSON'],PALETTE['violet']),
      ('AlertRule','AlertRule 预警规则',['id/eventType: UUID/String','condition: DSL/JSON','severity/slaMinutes: String/Int','routingPolicy: JSON'],'white'),
      ('Alert','AlertCase 例外工单',['id/ruleId: UUID','businessRef/severity/status: String','ownerId: UUID?','detected/resolvedAt: Instant','rootCause/resolution: JSON'],'white'),
      ('Notification','Notification 通知',['id/alertCaseId: UUID?','channel/recipient/templateId: String','status: SendStatus'],'white'),
      ('Metric','MetricDefinition 指标',['id/code/name: UUID/String','formula: DSL/SQL','dimensions/refreshPolicy: JSON/String'],'white'),
      ('Snapshot','MetricSnapshot',['id/metricId: UUID','timeBucket: Instant','dimensions: JSON','value: Decimal'],'white'),
      ('Workflow','WorkflowDefinition 工作流',['id/code/version: UUID/String','businessType: String','definition: BPMN/JSON','status: DefinitionStatus'],'white'),
      ('Instance','WorkflowInstance',['id/definitionId: UUID','businessRef/currentNode: String','status: WorkflowStatus','variables: JSON'],'white'),
      ('Task','WorkflowTask',['id/instanceId: UUID','nodeKey/assigneeId: String','dueAt: Instant','status: TaskStatus'],'white'),
      ('Endpoint','IntegrationEndpoint 集成端点',['id/code/name: UUID/String','type: REST/EDI/SFTP/WEBHOOK','authConfigRef/mappingVersion: String','status: Status'],'white'),
      ('Message','IntegrationMessage 集成消息',['id/endpointId: UUID','direction: IN/OUT','businessKey/correlationId: String','payloadObjectId: UUID','status/retryCount: Enum/Int'],'white')]
    for c in classes: dot+=class_node(*c)
    for a,b,l in [('Rule','Calc','0..*'),('Calc','VLine','0..1'),('Voucher','VLine','1..*'),('Voucher','Recon','0..1'),('Recon','Invoice','0..*'),('Invoice','Payment','0..*'),('Event','AlertRule','matches'),('AlertRule','Alert','0..*'),('Alert','Notification','0..*'),('Metric','Snapshot','0..*'),('Workflow','Instance','0..*'),('Instance','Task','1..*'),('Endpoint','Message','0..*')]: dot+=f'{a} -> {b} [label="{l}", arrowhead=none];\n'
    dot+='}'; out['47_class_billing_control']=render_dot('47_class_billing_control',dot)
    return out

# --------------------------- DOCX helpers ---------------------------

def set_cell_shading(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=70, start=70, bottom=70, end=70):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None: node=OxmlElement(f'w:{m}'); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')


def set_row_cant_split(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:cantSplit'); trPr.append(el)


def set_repeat_header(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)


def set_run_font(run, size=9, bold=False, color='1F2937', italic=False, name='Noto Sans CJK SC'):
    run.font.name=name; run._element.rPr.rFonts.set(qn('w:eastAsia'),name)
    run.font.size=Pt(size); run.bold=bold; run.italic=italic; run.font.color.rgb=RGBColor.from_string(color)


def keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next=value


def add_body(doc: Document, text: str, bold_prefix: Optional[str]=None):
    p=doc.add_paragraph(style='Body Text')
    if bold_prefix and text.startswith(bold_prefix):
        r=p.add_run(bold_prefix); set_run_font(r,9.1,True)
        r=p.add_run(text[len(bold_prefix):]); set_run_font(r,9.1)
    else:
        r=p.add_run(text); set_run_font(r,9.1)
    p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.12
    return p


def add_bullets(doc: Document, items: Iterable[str]):
    for item in items:
        p=doc.add_paragraph(style='List Bullet'); set_run_font(p.add_run(str(item)),8.9)
        p.paragraph_format.space_after=Pt(2.5); p.paragraph_format.line_spacing=1.08


def add_callout(doc: Document, title: str, text: str, kind='info'):
    colors={'info':('EAF2FF','2A6FDB'),'warn':('FFF4D6','C56A00'),'success':('E8F5E9','087F8C'),'danger':('FDECEC','B42318')}
    fill,border=colors.get(kind,colors['info'])
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=True
    c=t.cell(0,0); set_cell_shading(c,fill); set_cell_margins(c,110,140,110,140)
    tcPr=c._tc.get_or_add_tcPr(); borders=tcPr.find(qn('w:tcBorders'))
    if borders is None: borders=OxmlElement('w:tcBorders'); tcPr.append(borders)
    for side in ('top','left','bottom','right'):
        el=OxmlElement(f'w:{side}'); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'8'); el.set(qn('w:color'),border); borders.append(el)
    p=c.paragraphs[0]; r=p.add_run(title+'  '); set_run_font(r,9.3,True,border)
    r=p.add_run(text); set_run_font(r,8.9,False,'1F2937')
    doc.add_paragraph().paragraph_format.space_after=Pt(1)


def add_table(doc: Document, headers: List[str], rows: List[List[str]], widths: Optional[List[float]]=None, font_size=7.8, header_fill='16324F'):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    hdr=t.rows[0]; set_repeat_header(hdr); set_row_cant_split(hdr)
    for i,h in enumerate(headers):
        c=hdr.cells[i]; c.text=''; set_cell_shading(c,header_fill); set_cell_margins(c,55,60,55,60); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0); set_run_font(p.add_run(str(h)),font_size,True,'FFFFFF')
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells; set_row_cant_split(t.rows[-1])
        for i,val in enumerate(row):
            c=cells[i]; c.text=''; set_cell_margins(c,45,55,45,55); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            if ridx%2: set_cell_shading(c,'F7F9FB')
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
            set_run_font(p.add_run(str(val)),font_size)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Cm(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t


def add_figure(doc: Document, path: Path, caption: str, width_cm: float, max_height_cm: Optional[float]=None):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(1); p.paragraph_format.keep_with_next=True
    r=p.add_run()
    with Image.open(path) as im: pxw,pxh=im.size
    h=width_cm*pxh/pxw
    if max_height_cm is not None and h>max_height_cm:
        shape=r.add_picture(str(path),height=Cm(max_height_cm))
    else:
        shape=r.add_picture(str(path),width=Cm(width_cm))
    # Accessibility: expose the human-readable figure caption as image alt text.
    shape._inline.docPr.set('descr', caption)
    shape._inline.docPr.set('title', caption)
    cp=doc.add_paragraph(style='Figure Caption'); cp.paragraph_format.keep_together=True; cp.paragraph_format.space_after=Pt(5)
    set_run_font(cp.add_run(caption),8.0,False,'5B6472',True)
    return cp


def add_code(doc: Document, text: str):
    t=doc.add_table(rows=1,cols=1); t.style='Table Grid'; c=t.cell(0,0); set_cell_shading(c,'F3F5F7'); set_cell_margins(c,90,110,90,110)
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    for idx,line in enumerate(text.strip('\n').splitlines()):
        if idx: p.add_run('\n')
        set_run_font(p.add_run(line),7.7,False,'1F2937',False,'DejaVu Sans Mono')
    doc.add_paragraph().paragraph_format.space_after=Pt(1)


def set_landscape(section, landscape=True):
    if landscape:
        section.orientation=WD_ORIENT.LANDSCAPE; section.page_width=Cm(29.7); section.page_height=Cm(21)
        section.left_margin=Cm(1.35); section.right_margin=Cm(1.35); section.top_margin=Cm(1.25); section.bottom_margin=Cm(1.25)
    else:
        section.orientation=WD_ORIENT.PORTRAIT; section.page_width=Cm(21); section.page_height=Cm(29.7)
        section.left_margin=Cm(1.65); section.right_margin=Cm(1.65); section.top_margin=Cm(1.55); section.bottom_margin=Cm(1.55)


def new_section(doc: Document, landscape=False):
    sec=doc.add_section(WD_SECTION.NEW_PAGE); set_landscape(sec,landscape); sec.header.is_linked_to_previous=True; sec.footer.is_linked_to_previous=True; return sec


def add_section_title(doc: Document, number: str, title: str, subtitle: Optional[str]=None):
    p=doc.add_heading(f'{number}  {title}',level=1); keep_with_next(p,True)
    if subtitle:
        p=doc.add_paragraph(); set_run_font(p.add_run(subtitle),8.7,False,'5B6472',True); p.paragraph_format.space_after=Pt(7)


def configure_doc(doc: Document):
    sec=doc.sections[0]; set_landscape(sec,False)
    styles=doc.styles
    normal=styles['Normal']; normal.font.name='Noto Sans CJK SC'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC'); normal.font.size=Pt(9)
    body=styles['Body Text']; body.font.name='Noto Sans CJK SC'; body._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC'); body.font.size=Pt(9.1)
    for level,size,color in [(1,17,'16324F'),(2,12.5,'2A6FDB'),(3,10.5,'087F8C')]:
        st=styles[f'Heading {level}']; st.font.name='Noto Sans CJK SC'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC'); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
        st.paragraph_format.space_before=Pt(10 if level==1 else 7); st.paragraph_format.space_after=Pt(5); st.paragraph_format.keep_with_next=True
    if 'Figure Caption' not in styles:
        st=styles.add_style('Figure Caption',WD_STYLE_TYPE.PARAGRAPH)
    else: st=styles['Figure Caption']
    st.font.name='Noto Sans CJK SC'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC'); st.font.size=Pt(8); st.font.italic=True; st.font.color.rgb=RGBColor.from_string('5B6472'); st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    # Header/footer for all sections. These remain linked.
    hp=sec.header.paragraphs[0]; hp.text=''; set_run_font(hp.add_run('SCM Cloud 同类系统 - 架构与功能设计说明书  |  v0.1'),7.5,True,'5B6472')
    fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=fp.add_run(); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); run._r.addnext(fld)


# --------------------------- Detailed feature model ---------------------------

def F(code,name,actors,behavior,io,evidence='C'):
    return [code,f'{name}\n参与者：{actors}',behavior,io,evidence]

FEATURE_SECTIONS: List[Tuple[str,str,List[List[str]]]]=[]
platform_features = [
F('PLT-001','租户开通与生命周期','平台管理员、租户管理员','创建租户代码、默认语言/时区/币种、套餐能力和初始管理员；支持启用、暂停、到期和归档。所有业务表必须携带 tenant_id，租户切换必须重新计算权限上下文。','输入：租户资料、隔离模式；输出：Tenant、默认组织、管理员。状态：Provisioning→Active→Suspended→Archived。','C'),
F('PLT-002','多组织与业务域树','租户管理员','维护集团、法人、事业部、仓库和运输组织树；组织路径用于权限继承、审批路由、配置覆盖和汇总。移动节点时做循环校验并异步重建 path。','输入：组织类型、父组织；输出：组织树、数据范围。','A/B [S01][S08]'),
F('PLT-003','用户、人员与账号分离','租户管理员','账号负责认证，人员档案负责业务责任；支持批量导入、启停、解锁、有效期和离职交接，历史单据保留责任人快照。','输出：User、Person、IdentityBinding。异常：重复账号、身份源冲突。','B/C'),
F('PLT-004','SSO、MFA 与 API 凭证','安全管理员、集成管理员','支持 OIDC/SAML/企业目录、本地账号、多因素和服务账号；令牌包含租户、组织、设备与权限版本，敏感操作可二次认证。','输出：Session、Token、ApiCredential；审计登录成功/失败和因子变更。','C'),
F('PLT-005','角色与资源权限','租户管理员、安全管理员','资源粒度覆盖菜单、页面、API、按钮、字段和导出；角色可继承模板，后端命令必须再次鉴权，前端隐藏不作为安全边界。','输入：Role、Permission；输出：可审计 Allow/Deny 决策。','A/B [S08]'),
F('PLT-006','数据权限与 ABAC','租户管理员','按组织树、仓库、货主、伙伴、创建者、区域或自定义属性限定数据；策略采用显式表达式并缓存，缓存键包含权限版本。','输出：DataScope、PolicyDecision；异常：范围表达式无效。','C'),
F('PLT-007','工作台与多标签会话','全体 Web 用户','左侧模块导航、顶部租户/组织/仓库切换、工作区多标签、最近访问和收藏；标签保存查询条件与未提交草稿，关闭前提示脏数据。','输出：WorkspaceLayout、PageSession、Favorite。','B [S08]'),
F('PLT-008','待办与消息中心','全体用户','聚合审批、异常、到期任务、接口失败和通知；支持按严重度、业务域、责任组筛选，一键跳转到业务对象并记录已读。','输出：InboxItem、NotificationRead。','B/C'),
F('PLT-009','分层配置与版本发布','实施顾问、租户管理员','按租户→组织→仓库→客户层级继承与覆盖；发布前做差异预览、依赖校验、灰度生效和回滚。业务对象记录实际使用的配置版本。','输出：ConfigVersion、PublishRecord。','C'),
F('PLT-010','业务字典与原因码','主数据管理员','维护状态原因、异常类型、单位、包装级别、温层、服务类型等；停用字典不能影响历史快照，关键原因码要求备注或附件。','输出：Dictionary、DictionaryItem。','B/C'),
F('PLT-011','单号与编码规则','实施顾问','按租户、组织、业务类型、日期和流水生成订单号、任务号、LPN、运单号和预约号；高并发使用号段或数据库序列确保不重复。','输出：NumberRule、SequenceReservation。','C'),
F('PLT-012','工作流设计与版本化','业务管理员、实施顾问','以节点、条件、候选人、超时和回退定义审批/任务流程；已运行实例绑定不可变版本，新版本只影响后续实例。','输出：WorkflowDefinition、WorkflowInstance。','B/C'),
F('PLT-013','统一审批中心','审批人、业务主管','支持同意、拒绝、退回、加签、转交、撤回和批量审批；候选人由组织、角色、金额和业务属性计算，动作写入审计。','状态：Pending→Approved/Rejected/Returned/Cancelled。','B/C'),
F('PLT-014','规则引擎与决策追踪','实施顾问、运营人员','承载订单分配、库位、波次、承运商、时隙和计费规则；按优先级求值，保存命中、排除和输入快照，支持模拟而不落单。','输出：RuleDecision、EvaluationTrace。','C'),
F('PLT-015','调度任务与异步作业','运维、业务管理员','支持一次性/周期/事件触发任务，包含并发限制、超时、重试、取消、进度和结果文件；多实例通过租约避免重复执行。','输出：JobDefinition、JobRun、JobLog。','C'),
F('PLT-016','附件与对象存储','全体用户','附件采用预签名上传，元数据与业务对象关联；支持病毒扫描、类型/大小限制、版本、保留期和权限校验，敏感附件水印下载。','输出：FileObject、AttachmentLink。','B/C'),
F('PLT-017','打印、标签与模板','仓库、运输、客服','模板按租户、客户、仓库、单据类型和语言版本化；支持箱唛、托盘标、拣货单、装箱单、运单、回单和预约凭证，打印任务可路由指定打印机。','输出：PrintTemplate、PrintJob、LabelData。','A/B [S03][S08][S09]'),
F('PLT-018','评论、@提及与协同记录','客服、计划员、伙伴','业务对象下支持评论、内部/外部可见性、@责任人、附件和解决标记；评论不能代替状态变更，但可成为例外工单证据。','输出：CommentThread、Mention、Activity。','B/C'),
F('PLT-019','通知模板与多渠道发送','系统、业务管理员','事件触发站内信、邮件、短信、微信或推送；模板变量白名单化，按订阅偏好、免打扰和升级策略路由，失败支持重试/降级。','输出：Notification、DeliveryAttempt。','A/B [S06]'),
F('PLT-020','审计日志与数据变更历史','安全、内审、管理员','记录登录、查询敏感数据、导出、审批、接口调用和业务变更；重要变更保存 before/after、traceId、IP 和设备，审计不可由业务用户删除。','输出：AuditLog、ChangeHistory。','C'),
F('PLT-021','批量导入、校验与回执','主数据、客服、计划员','上传 Excel/CSV 后先预检字段、字典、重复和业务规则；按行返回错误，不允许部分覆盖未知数据；大批量异步处理并生成回执。','状态：Uploaded→Validating→Ready→Importing→Completed/Failed。','B [S08]'),
F('PLT-022','导出与数据脱敏','全体用户','导出复用当前筛选、排序和列配置；按数据权限与字段策略脱敏，超阈值转异步任务并限制下载有效期，导出动作强制审计。','输出：ExportJob、DownloadToken。','B/C'),
F('PLT-023','统一搜索与保存视图','全体用户','支持业务号、外部号、伙伴、商品、状态和时间范围组合检索；保存个人/共享视图、列宽、排序、聚合和快速筛选。','输出：SavedView、SearchQuery。','B [S08]'),
F('PLT-024','国际化、时区与单位','跨区域用户','界面、字典、模板和通知支持多语言；时间以 UTC 存储、按业务时区显示；数量同时保存原单位和基础单位，金额包含币种与舍入规则。','输出：LocaleContext、UnitConversion。','C'),
F('PLT-025','特性开关与渐进发布','产品管理员、运维','按租户、组织、角色或百分比启用新功能；支持新旧流程并行、快速回退和实验指标，关键业务记录开关版本。','输出：FeatureFlagDecision。','C'),
]

mdm_features = [
F('MDM-001','商品与 SKU 主数据','主数据管理员','维护 SKU、名称、分类、基础单位、温层、危险属性、条码、批次/序列规则和保质期；变更影响新业务，历史单据使用快照。','输出：Product、ProductVersion、Barcode。','A/B [S02][S03]'),
F('MDM-002','包装层级与单位换算','主数据管理员','定义件、箱、托、重量和体积换算，支持非整倍包装及客户专用包装；任务生成引用当时的 PackageSpec 版本。','输出：PackageSpec、UomConversion。','A/B [S03][S09]'),
F('MDM-003','客户、供应商与承运商','主数据管理员、采购、运输','统一伙伴档案、角色、联系人、地址、税务、信用、证照和服务能力；同一法人可承担多个角色，外部编码通过映射表管理。','输出：Partner、PartnerRole、ExternalCodeMap。','A/B [S02][S04]'),
F('MDM-004','地址、地理编码与服务区域','主数据、运输计划员','地址标准化、经纬度、行政区、时间窗、进出场说明和地理围栏；地理编码失败进入人工校正，不覆盖原始文本。','输出：Address、GeoPoint、ServiceZone。','B/C'),
F('MDM-005','仓库、库区、库位与月台','仓库管理员','维护仓库、库区、通道、库位、暂存区、门岗、月台及容量；状态停用前检查库存、未完任务和未来预约。','输出：Warehouse、Location、Dock、Gate。','A/B [S03][S05]'),
F('MDM-006','车辆、设备与司机档案','运输管理员','维护车型、载重、容积、温控、车牌、司机证照和可用状态；证照到期触发提醒并阻止指派。','输出：Vehicle、EquipmentType、Driver、Certificate。','A/B [S04]'),
F('MDM-007','合同、费率卡与有效期','商务、财务','合同关联客户/供应商/承运商，费率按路线、区域、车型、重量、体积、里程、时间和附加服务版本化；重叠有效期必须冲突校验。','输出：Contract、RateCard、RateVersion。','A/B [S02][S04]'),
F('MDM-008','日历、班次与工作时间','仓库、运输、预约管理员','定义营业日、节假日、班次、截单时间和资源工作时段；计划、预约和 SLA 统一引用日历服务。','输出：BusinessCalendar、Shift、WorkingWindow。','B/C'),
F('MDM-009','主数据映射与外部编码','集成管理员','维护 ERP、客户、承运商和设备系统的外部编码映射；同一来源+对象类型+外部码唯一，变更保留版本。','输出：ExternalCodeMap、MappingVersion。','B/C'),
F('MDM-010','主数据质量与审批','主数据管理员、审批人','必填、重复、地址、证照、单位和引用完整性评分；低质量数据可草稿保存但不能参与履约，关键变更走审批。','输出：DataQualityIssue、ApprovalTask。','C'),
]
FEATURE_SECTIONS.append(('平台、租户与共性能力','多租户 SaaS 的身份、配置、规则、工作流、审计和统一工作台。',platform_features))
FEATURE_SECTIONS.append(('主数据与基础配置','OMS、WMS、TMS、AMS 和计费共享的受控业务语义。',mdm_features))
oms_features = [
F('OMS-001','多渠道订单接入','ERP、电商、客户、客服','支持 API、EDI、文件、门户和人工录入销售、采购、调拨、退货等订单；统一转换为内部订单模型，保留来源报文和映射版本。','输出：BusinessOrder、RawMessageRef。','A [S02]'),
F('OMS-002','幂等去重与外部编号','集成系统、客服','按租户+渠道+外部订单号+版本或 Idempotency-Key 去重；重复内容返回原结果，内容冲突创建数据异常而非覆盖。','输出：IdempotencyRecord、DuplicateCase。','B/C'),
F('OMS-003','订单字段与业务规则校验','系统、客服','校验伙伴、商品、地址、数量、单位、时间窗、信用、禁运和必填扩展字段；错误按字段返回，警告可经授权强制通过。','状态：Draft/Invalid→Open。','A/B [S02]'),
F('OMS-004','订单草稿与版本控制','客服、客户','草稿可保存未完整数据；正式提交后使用乐观锁和变更版本，外部变更必须引用原版本，避免覆盖并发处理。','输出：OrderVersion、ChangeSet。','B/C'),
F('OMS-005','订单审核与风险检查','审批人、客服','根据金额、客户、品类、信用和异常标记决定自动审核或人工审批；审核通过后冻结关键商业字段，后续修改走变更单。','状态：Open→Approved/Rejected/Hold。','A/B [S02]'),
F('OMS-006','合单与拆单','客服、系统','按客户、地址、时间窗、仓库、温层和运输模式合并/拆分；保留原订单与履约单映射，金额和数量分摊可追溯。','输出：MergeGroup、SplitRelation。','A [S02]'),
F('OMS-007','优先级与履约排序','客服、运营','支持订单级和行级优先级、VIP、加急、截单时间和 SLA 权重；优先级变化触发重新分配但不得抢占已实物执行库存。','输出：PriorityDecision。','A/B [S02]'),
F('OMS-008','订单冻结与解冻','客服、风控、财务','按支付、信用、地址、合规、客户请求或异常冻结；冻结可限定订单或行，阻止释放但保留分配；解冻要求原因和权限。','输出：OrderHold、ReleaseRecord。','A [S02]'),
F('OMS-009','库存可承诺 ATP 查询','客服、客户','聚合 WMS 可用库存、在途、预期入库和安全库存，按仓库/货主/批次返回可承诺日期；结果带快照时间和不确定性。','输出：AvailabilityPromise。','A/B [S02]'),
F('OMS-010','履约来源与仓库分配','系统、运营','按库存、距离、服务区、成本、产能、时限、客户偏好和分单上限选择仓库/供应商；保存候选、排除原因和规则版本。','输出：Allocation、SourcingDecision。','A [S02]'),
F('OMS-011','分配预占与释放','系统、库存计划员','分配成功后向 WMS 预占或创建履约请求；超时、取消或策略变更释放预占，跨域通过幂等命令和事件确认。','状态：Proposed→Reserved→Released/Failed。','A/B [S02]'),
F('OMS-012','履约单生成','系统、仓库','按仓库、货主、订单类型生成入库、出库、调拨履约单和行；履约单独立状态机，OMS 通过事件聚合进度。','输出：FulfillmentOrder、FulfillmentLine。','A [S02]'),
F('OMS-013','运输需求生成','系统、运输计划员','基于提货/送货地址、时间窗、重量体积、温层和服务等级生成运输需求；直运和多段运输可拆分多个请求。','输出：ShipmentRequest。','A [S02][S04]'),
F('OMS-014','订单释放与截单','客服、系统','审核、分配和依赖满足后释放；支持批量、按日历自动释放和截单前修改，释放后变更需检查仓储/运输执行进度。','状态：Approved/Allocated→Released。','A [S02]'),
F('OMS-015','伙伴订单协同','客户、供应商、客服','客户可确认、修改或取消请求；供应商可接单、承诺交期、反馈缺货和发货信息；每次协同形成结构化事件和评论。','输出：PartnerConfirmation、PromiseDate。','A [S02]'),
F('OMS-016','供应商 ASN 预告','供应商、采购、仓库','供应商按采购订单创建 ASN，包含箱/托、批次、预计到达和运输信息；超量/重复/过期校验后推送 WMS 与 AMS。','输出：PartnerASN、ASNLine。','A/B [S02][S03]'),
F('OMS-017','订单状态时间线','全体业务用户','将订单、履约、仓库、运输、预约、结算事件按业务时区汇总；显示来源、操作者、原状态、新状态、附件和 traceId。','输出：OrderTimelineProjection。','A/B [S02]'),
F('OMS-018','订单变更单','客户、客服、审批人','正式订单修改数量、地址、日期、服务级别时创建变更单；评估库存、任务、运单和费用影响，需各域确认后原子生效或补偿。','输出：OrderChange、ImpactAssessment。','B/C'),
F('OMS-019','订单取消与撤销执行','客户、客服','按状态判断可直接取消、需撤销预占、需取消仓库任务或需运输改派；不可逆执行后转退货/补偿流程。','状态：Open/Approved/Allocated→Cancelled。','A/B [S02]'),
F('OMS-020','部分履约与欠货','客服、库存计划员','订单行可分批履约，记录已承诺、已分配、已发运、已交付和取消数量；欠货可等待、替代、改仓或取消。','输出：Backorder、PartialFulfillment。','A/B [S02]'),
F('OMS-021','替代品与客户确认','客服、客户','按商品替代组、价格和兼容约束推荐替代品；需要客户确认的替代在时限内未响应则按策略取消或继续等待。','输出：SubstitutionProposal、Decision。','B/C'),
F('OMS-022','退货授权 RMA','客户、客服','校验原订单、可退数量、时间窗、商品条件和费用；生成 RMA、逆向运输或退仓请求，支持退款、换货、维修和拒绝。','状态：Requested→Authorized→Received→Resolved→Closed。','A/B [S02]'),
F('OMS-023','订单异常与人工工作台','客服、运营','聚合校验失败、分配失败、长时间未释放、履约短缺、运输延误和签收差异；可批量指派、冻结、重试或创建补偿。','输出：OrderExceptionCase。','B/C'),
F('OMS-024','SLA 与承诺时间监控','客服、控制塔','计算确认、分配、释放、发货和交付 SLA，考虑日历和暂停原因；临近超时预警，超时升级并保存责任域。','输出：SlaClock、SlaBreachEvent。','B/C'),
F('OMS-025','订单费用与结算请求','客服、财务','订单商业费用、折扣、附加服务和履约事实汇总为结算请求；实际仓储/运输费用由计费域计算，不在 OMS 直接改账。','输出：SettlementRequest、ChargeFactRef。','A/B [S02]'),
F('OMS-026','订单查询、报表与批量动作','客服、运营','支持多条件查询、保存视图、列聚合、批量审核/冻结/释放/导出；批量动作逐单鉴权和校验，返回成功/失败明细。','输出：BatchCommandResult。','B [S08]'),
F('OMS-027','客户门户自助服务','客户','客户查看订单、承诺、库存可用、发货、轨迹、签收、附件和对账状态；可在权限和状态允许时修改、取消、创建退货。','输出：CustomerOrderView。','A/B [S01][S02]'),
]
FEATURE_SECTIONS.append(('OMS 订单管理与履约编排','订单从多渠道进入，到分配、释放、仓储/运输执行和客户协同的全生命周期。',oms_features))
wms_features = [
F('WMS-IN-001','入库单与预期收货','采购、仓库、系统','接收采购、退货、调拨、生产或 ASN 生成入库单；按仓库、货主和来源拆分，校验商品、包装、批次和预计到达。','状态：Draft→Expected→Arrived→Receiving→Completed。','A [S03][S09]'),
F('WMS-IN-002','ASN 与包装层级解析','供应商、收货员','解析箱、托、LPN、商品和数量层级；支持完整 ASN、简化 ASN 和无 ASN 收货，重复 LPN 或超量按策略拦截。','输出：InboundPackage、HandlingUnit。','A [S03][S09]'),
F('WMS-IN-003','入库预约关联','预约管理员、仓库','入库单/ASN 可创建或关联 AMS 预约；预约确认回写预计到场、车辆和月台，改期不改变订单数量。','输出：AppointmentOrderLink。','A [S05][S10]'),
F('WMS-IN-004','到场签到与月台接收','门岗、收货员','门岗签到后更新到场事件；按预约、优先级和月台状态开启收货，未预约车辆可走临时登记和审批。','输出：ArrivalEvent、DockAssignmentRef。','A/B [S05]'),
F('WMS-IN-005','收货任务创建与指派','收货主管','按入库单、月台、工作量和班组生成任务；支持自动派工、抢单、转派、暂停和并行收货，任务版本防止多人重复。','输出：ReceiptTask、LaborAssignment。','A/B [S03]'),
F('WMS-IN-006','条码扫描与商品识别','收货员、RF','扫描订单、商品、箱码、托盘码和库位；支持 GS1/客户码映射，无法识别进入人工选择并记录原始条码。','输出：ScanEvent、BarcodeResolution。','A [S03][S09]'),
F('WMS-IN-007','盲收与按单收货','收货员','可配置隐藏预期数量的盲收，或显示订单数量的按单收货；超收、短收和替代包装需要原因与授权。','输出：ReceiptLine、ReceivingVariance。','A/B [S03]'),
F('WMS-IN-008','批次、序列号与效期采集','收货员','按商品规则强制采集供应商批次、生产/失效日期和序列号；序列号租户范围唯一，效期不满足最短剩余天数时隔离。','输出：InventoryLot、SerialNumber。','A [S03][S09]'),
F('WMS-IN-009','LPN 建托、拆托与合托','收货员','生成箱/托 LPN 及父子关系，支持混托限制、拆分和合并；LPN 变更产生处理单元事件，标签可补打但编号不重复。','输出：HandlingUnit、LabelJob。','A [S03][S09]'),
F('WMS-IN-010','收货差异与拒收','收货员、主管','记录数量、包装、破损、温度和文件差异；可补收、拒收、待定或转质检，差异通知供应商/采购并保留照片。','输出：ReceivingException、Disposition。','A/B [S03]'),
F('WMS-IN-011','质检计划与抽样','质检员','按商品、供应商、批次和风险选择免检、全检或抽样计划；记录检验项目、样本、结果和附件，未放行库存不可用。','状态：Pending→Inspecting→Accepted/Rejected/Hold。','A [S03][S09]'),
F('WMS-IN-012','不合格品处置','质检、采购、仓库','不合格库存可退供应商、返工、降级、报废或特采放行；处置需要权限和原因，变更库存状态并产生费用事实。','输出：QualityDisposition、InventoryMovement。','A/B [S03]'),
F('WMS-IN-013','上架策略','仓库管理员、系统','按固定/动态库位、同品同批、温层、危险属性、容量、周转和拣选补货需求推荐库位；保存候选与排除原因。','输出：PutawayDecision、StrategyTrace。','A [S03][S09]'),
F('WMS-IN-014','上架任务与路径','上架员、RF','按 LPN/商品生成上架任务，支持整托、拆零和多库位；推荐任务顺序，扫描源/目标库位和 LPN 后确认移动。','输出：PutawayTask、Movement。','A [S03]'),
F('WMS-IN-015','越库入库','仓库、系统','匹配紧急出库需求时将收货分配到越库区，不进入普通存储；需要商品、批次、数量、质检和时间窗满足约束。','输出：CrossDockAllocation。','A [S09][S10]'),
F('WMS-IN-016','入库完成与回写','系统','所有收货、质检、上架或处置完成后关闭入库单；发布库存可用和入库完成事件，回写 OMS/ERP，不允许仅靠页面标记完成。','输出：inbound.completed.v1。','A/B [S03]'),

F('WMS-INV-001','库存余额与维度','库存管理员','库存按租户、仓库、库位、货主、商品、批次、序列、LPN 和状态聚合；保存现存、可用、分配、冻结、在途数量和版本。','输出：InventoryBalance。','A [S03]'),
F('WMS-INV-002','库存流水与可追溯性','库存、内审','每次收货、上架、拣选、移动、调整和发运写不可变流水，包含业务引用、前后维度、数量、操作者和 traceId。','输出：InventoryMovement、TraceChain。','A/B [S03]'),
F('WMS-INV-003','库存状态与质量状态','库存、质检','支持可用、待检、冻结、破损、过期、待处置等状态；状态转换由受控命令执行，不通过直接修改余额。','输出：InventoryStatusChange。','A/B [S03]'),
F('WMS-INV-004','冻结与解冻','库存主管、客服','按订单、批次、LPN、库位或数量冻结；冻结减少可用但不改现存，解冻检查未完成处置和审批。','输出：InventoryHold。','A/B [S03]'),
F('WMS-INV-005','库存预占与释放','OMS、WMS','为订单/波次预占库存，使用乐观锁或原子 SQL 防超卖；取消、超时和短拣释放，预占不能超过可用量。','输出：Reservation、AllocationDetail。','A/B [S02][S03]'),
F('WMS-INV-006','库内移库','库存员、RF','从源库位/LPN 移到目标库位，支持整托和部分数量；先校验源库存与目标容量，确认时一次事务写余额和流水。','输出：TransferTask、Movement。','A [S03]'),
F('WMS-INV-007','货主转换','库存主管、财务','在合同和授权允许时转换库存所有权；商品、批次、数量不变，生成双向流水和可能的结算事实。','输出：OwnershipTransfer。','B/C'),
F('WMS-INV-008','库位容量与混放约束','仓库管理员','按数量、重量、体积、托位和属性限制容量；危险品、温层、批次/货主混放规则在上架和移库时统一校验。','输出：CapacityCheck。','A/B [S03]'),
F('WMS-INV-009','循环盘点','库存主管、盘点员','按 ABC、差异风险、移动频次或日程生成循环盘点；可盲盘、复盘，盘点范围按策略冻结或允许带事务盘点。','输出：CountOrder、CountLine。','A [S03]'),
F('WMS-INV-010','全库盘点','仓库主管、财务','按仓库/货主冻结范围，生成盘点批次和区域任务；支持初盘、复盘、差异审批、分段解冻和期间报告。','状态：Planned→Counting→Reviewing→Posted→Closed。','A [S03]'),
F('WMS-INV-011','库存调整','库存主管、审批人','针对盘盈、盘亏、破损、数据修复创建调整单；必须选择原因、附件和财务影响，审核后写流水。','输出：AdjustmentOrder。','A/B [S03]'),
F('WMS-INV-012','自动补货','系统、补货员','按拣选位最小/最大库存、波次需求和前置时间生成补货任务；源库存按 FEFO/FIFO 选择，避免与盘点/冻结冲突。','输出：ReplenishmentTask。','A [S03][S09]'),
F('WMS-INV-013','库存老化与效期预警','库存、客户','按入库日期、生产日期和失效日期计算库龄；临期、过期和长期滞销触发预警、冻结或促销/退货建议。','输出：AgingBucket、ExpiryAlert。','A/B [S03]'),
F('WMS-INV-014','批次与序列追溯','质量、客服','从供应商/收货批次追溯到库位、拣选、包裹、客户和签收；反向从客户序列号定位采购和质检。','输出：GenealogyGraph、RecallList。','A/B [S03]'),
F('WMS-INV-015','库存查询与快照','客服、库存、客户','按多维度查询现存、可用、预占、冻结和在途；读模型标注快照时间，关键承诺需再走强一致预占命令。','输出：InventoryView、SnapshotAt。','A [S03]'),
F('WMS-INV-016','库存对账与 ERP 同步','库存、财务、集成','按日/期间汇总期初、收、发、调、盘、期末并与 ERP 对账；差异进入工单，不允许无审计地改平。','输出：InventoryReconciliation。','B/C'),

F('WMS-OUT-001','出库单接收与释放','OMS、仓库','接收销售、调拨、退供应商等出库履约单；校验商品、地址、服务级别和截止时间，释放后进入波次或单单作业。','状态：Draft→Released。','A [S03]'),
F('WMS-OUT-002','波次模板与选择器','仓库管理员','按客户、承运方式、路线、温层、订单类型、截止时间、商品和库区定义波次模板；支持模拟订单数、行数和工作量。','输出：WaveTemplate、WaveCandidate。','A [S03][S09]'),
F('WMS-OUT-003','波次计划与发布','波次计划员','手动或定时创建波次，计算产能和截止时间；发布前可调整订单，发布后生成分配/任务，取消需检查执行进度。','状态：Draft→Planned→Released→Completed。','A [S03]'),
F('WMS-OUT-004','库存分配策略','系统、库存计划员','按货主、状态、库区、整箱优先、FIFO/FEFO、批次、效期和最少拆分分配；并发更新余额版本防止超分配。','输出：AllocationDetail、StrategyTrace。','A [S03][S09]'),
F('WMS-OUT-005','缺货与重分配','库存计划员','分配不足时可等待入库、触发补货、替代批次、改仓、拆单或短发；处理结果回写 OMS 承诺。','输出：ShortageCase、Reallocation。','A/B [S03]'),
F('WMS-OUT-006','拣选任务拆分与合并','仓库主管','按区域、设备、容器、线路和工作量拆分任务；可按订单、批量、播种、分区、整箱/拆零模式组织，避免跨温区混合。','输出：PickTask、PickTaskLine。','A [S03][S09]'),
F('WMS-OUT-007','最短路径与任务排序','系统、拣选员','基于库位序、通道方向、拥堵和优先级生成路线；路线可动态调整但必须保留任务行和确认历史。','输出：PickRoute、RouteSeq。','A [S03]'),
F('WMS-OUT-008','RF 拣选扫描','拣选员、RF','按任务扫描源库位、商品/LPN、数量和目标容器；错误扫描即时阻止，连续扫描可离线排队并带设备序号。','输出：PickConfirmation、ScanEvent。','A [S03]'),
F('WMS-OUT-009','短拣与异常处理','拣选员、主管','短拣需原因码，可触发复核、库存冻结、循环盘点、重分配或短发；不能直接把任务标完成掩盖差异。','输出：ShortPickCase。','A/B [S03]'),
F('WMS-OUT-010','复核与差异纠正','复核员','按订单、容器或包裹扫描复核商品、数量、批次和序列号；错拣退回、补拣或重新分配，保留原确认。','输出：VerificationResult。','A [S03]'),
F('WMS-OUT-011','包装与箱型推荐','包装员、系统','根据商品尺寸、重量、禁配和客户规则推荐箱型；支持多箱、套箱、缓冲材料和包装服务，记录实际重量体积。','输出：PackTask、PackageUnit。','A/B [S03]'),
F('WMS-OUT-012','称重、量方与超差','包装员、设备','连接电子秤/量方设备，采集实际数据；与理论值偏差超阈值时阻止封箱并触发复核，设备读数保留来源。','输出：Measurement、WeightException。','A/B [S03]'),
F('WMS-OUT-013','出库标签与单据','包装、打印服务','按承运商、客户和目的地生成箱标、面单、装箱单和危险品文件；标签版本与包裹绑定，作废/补打需审计。','输出：ShippingLabel、PrintJob。','A [S03]'),
F('WMS-OUT-014','集货与暂存','集货员','按路线、运单、车次和装车顺序将包裹/LPN 移到暂存位；扫描防止混线路，暂存容量和超时受监控。','输出：StagingTask。','A/B [S03]'),
F('WMS-OUT-015','装车与装载校验','装车员、司机','按 TMS 运单/月台扫描车辆、包裹、托盘和顺序；校验应装/实装、重量体积、温层和封签，漏装或错装阻止发运。','输出：LoadTask、LoadConfirmation。','A/B [S03][S04]'),
F('WMS-OUT-016','发运确认与库存扣减','仓库主管','装车完成并确认发运后将库存从在库转为已发，关闭相关任务并发布 outbound.shipped 事件；失败可重试但必须幂等。','状态：Loaded→Shipped。','A [S03]'),
F('WMS-OUT-017','运输交接与可提货事件','系统、TMS','出库准备完成发布可提货事件，包含包裹、重量、体积、月台和时间；TMS 运单变更通过引用更新，不直接改 WMS 任务。','输出：outbound.ready.v1。','A/B [S03][S04]'),
F('WMS-OUT-018','出库取消与逆操作','客服、仓库','按执行阶段撤销波次、释放分配、取消任务或从暂存退回；已发运不可直接取消，转退货/召回流程。','输出：CancellationPlan、CompensationSteps。','B/C'),

F('WMS-VAS-001','贴标与换标','增值作业员','按客户/法规规则创建贴标任务，扫描原 LPN 和新标签；旧标签作废但历史可追溯，完成产生服务数量。','输出：ValueAddedOrder、ChargeFact。','A [S03]'),
F('WMS-VAS-002','重新包装与分装','增值作业员','支持拆箱、分装、组合、换箱和加固；输入/输出处理单元与数量守恒，材料消耗单独记录。','输出：RepackOperation、MaterialUsage。','A/B [S03]'),
F('WMS-VAS-003','组套与拆套','增值作业员','按 BOM/套装规则消耗组件生成套装 SKU，或反向拆套；批次/序列继承和成本分配可配置。','输出：KittingOrder、ComponentConsumption。','A/B [S03]'),
F('WMS-VAS-004','简单加工与装配','增值作业员、质检','支持装配、加赠品、说明书、刻码等步骤；工艺版本、工位、人员、开始/完成和质检结果可追溯。','输出：AssemblyTask、ProcessTrace。','A/B [S03]'),
F('WMS-LAB-001','劳务任务与标准工时','仓库主管','收货、上架、拣选、包装、装车和增值任务映射劳务类型与标准工时；支持班组、个人和外包队伍。','输出：LaborTask、StandardTime。','A [S03]'),
F('WMS-LAB-002','人员绩效与工作量','仓库主管、员工','按完成量、标准/实际工时、质量、异常和等待时间计算绩效；指标只用于改进，不覆盖原业务事实。','输出：LaborMetric、PerformanceScore。','A/B [S03]'),
F('WMS-DEV-001','RF 离线与同步','RF 用户、运维','短时离线缓存已授权任务和扫描命令，命令带设备序号、业务版本和幂等键；重连后顺序同步，冲突进入人工处理。','输出：OfflineCommandQueue、SyncConflict。','C'),
F('WMS-DEV-002','设备与自动化接口','仓库、集成','对接输送线、分拣机、称重、电子标签、打印机和门禁；设备消息统一转换为任务/事件，指令需回执和超时监控。','输出：DeviceCommand、DeviceEvent。','A/B [S03]'),
F('WMS-RPT-001','仓库作业看板','仓库主管、管理层','展示入库、出库、库存、任务积压、产能、人员绩效、月台和异常；指标可按仓库、货主、客户、班次下钻。','输出：WarehouseDashboard。','A/B [S03][S06]'),
F('WMS-BIL-001','仓储计费事实','WMS、财务','在收货、存储、操作、耗材和增值服务完成时产生可计费事实；事实包含数量、单位、业务引用和发生时间，不直接计算最终金额。','输出：WarehouseChargeFact。','A/B [S03]'),
]
FEATURE_SECTIONS.append(('WMS 仓储管理','覆盖入库、库存、出库、增值服务、劳务、RF 和设备集成。',wms_features))
tms_features = [
F('TMS-001','运输订单接收','OMS、客户、运输客服','接收销售、采购、调拨或独立运输需求，统一地址、时间窗、包装、重量体积、温层和服务等级；保留来源和外部编号。','状态：Open→Planned。','A [S04][S08]'),
F('TMS-002','运输订单审核','运输主管','校验地址、禁运、车辆约束、承运商资质、时间窗和费用责任；异常可冻结或退回，审核后进入计划。','输出：TransportOrderApproval。','A/B [S04]'),
F('TMS-003','计划批次与订单池','运输计划员','按日期、区域、模式、客户和优先级建立计划批次；订单池显示未计划、已锁定和冲突，计划员领取防止并发重复。','输出：PlanningBatch、PlanningLock。','A [S04][S08]'),
F('TMS-004','合单与拆运','运输计划员、系统','按提送地址、时间窗、车型、温层和服务规则合并订单，或拆分多车/多段；保留订单行与运单项映射。','输出：ConsolidationPlan、ShipmentItem。','A [S04]'),
F('TMS-005','运输模式与多段线路','运输计划员','支持公路、零担、快递、铁路、海空等模式及多段转运；每段有独立承运商、节点和 SLA，总订单汇总段状态。','输出：TransportLeg、ModeDecision。','A/B [S04]'),
F('TMS-006','车型与设备选择','系统、计划员','依据重量、体积、托位、温层、危险属性、装卸限制和客户要求选择车型/设备；超载或不兼容方案不可发布。','输出：EquipmentSelection。','A [S04][S08]'),
F('TMS-007','配载与装载率','计划员','将运输订单装入计划/运单，实时计算重量、体积、托位和价值利用率；允许手工拖放但重新校验约束。','输出：LoadPlan、UtilizationMetrics。','A [S04][S08]'),
F('TMS-008','路线与站点顺序','计划员、优化服务','根据提送节点、时间窗、里程、路况、车辆和司机工时生成路线；保存优化输入、约束、得分和手工调整。','输出：RoutePlan、StopSequence。','A/B [S04][S07]'),
F('TMS-009','路径优化与方案比较','计划员、管理层','生成成本、里程、准时率或车辆数不同权重的多个方案；用户可比较、锁定节点并重新优化，结果可解释。','输出：OptimizationScenario、ScoreBreakdown。','A/B [S07]'),
F('TMS-010','运力与容量管理','运输主管','维护自有/合同/临时运力、车型容量、区域、可用日历和预占；计划发布时原子占用容量，取消/拒单释放。','输出：CapacityReservation。','A [S04]'),
F('TMS-011','计划审批','运输主管、财务','按成本、偏差、承运商、加急和超载风险决定自动或人工审批；审批后关键计划字段锁定。','状态：Planned→Approved/Rejected。','A [S08]'),
F('TMS-012','直接委托承运商','计划员、承运商','选择合同承运商并发送运单、价格、时间窗和要求；承运商在有效期内接受、拒绝或提出问题。','输出：CarrierTender。','A [S04][S08]'),
F('TMS-013','竞价、询价与比价','采购、承运商','向候选承运商发布询价/竞价，记录报价、时效和附加条件；截止后按规则推荐但保留人工选择和审批。','输出：QuoteRequest、CarrierBid、AwardDecision。','A/B [S04]'),
F('TMS-014','承运商拒单与改派','承运商、计划员','拒绝、超时或资质失效时释放容量并重新计划；改派保留原委托记录、原因和价格差异，必要时触发升级。','输出：RetenderCase。','A [S08]'),
F('TMS-015','平台内转委托/分包','3PL、承运商','允许合规的下游合作方继续执行，但必须保留原始责任链、实际承运商、可见范围和费用分层。','输出：SubcontractAssignment。','A [S08]'),
F('TMS-016','车辆与司机指派','调度员、承运商','为运单指派车辆、司机和备用联系人；校验车型、证照、排班、冲突、温控和容量，变更通知相关方。','输出：VehicleAssignment。','A [S04][S08]'),
F('TMS-017','司机/车辆证照检查','调度员、司机','发运前检查驾驶证、行驶证、保险、危险品资质和车辆年检；到期或缺失阻止发运并生成例外。','输出：ComplianceCheck。','A/B [S04]'),
F('TMS-018','发运确认','调度员、司机、仓库','装车、封签、单据、车辆司机满足后确认发运；记录实际发车时间并进入跟踪，重复确认必须幂等。','状态：Accepted/Dispatched→Tracking。','A [S08]'),
F('TMS-019','运输节点与里程碑计划','计划员','为提货、到达、装卸、离开、交付和回单建立计划时间；节点可按运输模式和客户模板配置。','输出：Stop、MilestonePlan。','A [S04][S08]'),
F('TMS-020','司机 APP 节点上报','司机','司机接受任务、导航、签到、拍照、扫描、异常和签收；离线缓存节点，重连后按序同步，禁止跳过强制节点。','输出：DriverTask、TrackingEvent。','A [S01][S04]'),
F('TMS-021','GPS 与设备轨迹','设备、TMS','接收车载 GPS、手机定位或第三方 EDI 位置，清洗漂移和重复点；按车辆/运单关联并控制采样与保留期。','输出：PositionPoint、TrackSegment。','A [S04]'),
F('TMS-022','地理围栏自动事件','系统','车辆进入/离开提送货围栏自动生成到达/离开候选事件；与司机手工上报冲突时保留两者并按规则判定。','输出：GeofenceEvent。','A/B [S04]'),
F('TMS-023','ETA 预测与动态更新','控制塔、客户','基于计划、实时位置、历史、路况和站点作业更新 ETA；预测结果含置信度，显著变化触发通知和下游资源调整。','输出：EtaPrediction、EtaChangedEvent。','A/B [S04][S06]'),
F('TMS-024','在途可视化地图','运输客服、客户、管理层','地图显示车辆、线路、节点、当前位置、状态、ETA 和异常；按权限隐藏精确位置或司机个人信息。','输出：ShipmentMapProjection。','A [S04][S06]'),
F('TMS-025','运输异常检测','司机、系统、客服','识别延误、偏航、长停、温控、破损、拒收、证件和通信中断；按严重度创建工单、冻结费用或触发改派。','输出：TransportException。','A [S04][S06]'),
F('TMS-026','异常处置与升级','运输客服、承运商','记录责任人、处理计划、预计恢复、根因和证据；超 SLA 自动升级，关闭前验证业务状态和客户沟通。','输出：ExceptionAction、Escalation。','B/C'),
F('TMS-027','预约联动','TMS、AMS、司机','根据提送时间和场地要求创建/关联预约；预约改期更新站点计划和 ETA，运输取消释放时隙。','输出：ShipmentAppointmentLink。','A [S04][S05]'),
F('TMS-028','到达、签收与差异','司机、收货方','记录到达、开始/结束卸货、签收人、数量、拒收和货损；数量差异关联运单项，触发 OMS/WMS 处理。','输出：DeliveryConfirmation、DeliveryVariance。','A [S08]'),
F('TMS-029','POD 回单采集','司机、承运商','上传签名、照片、电子回单、盖章单据和备注；对象存储保存原件，元数据关联运单，支持补传与多页。','输出：ProofOfDelivery。','A [S01][S08]'),
F('TMS-030','POD 审核与退回','客服、财务','校验签收时间、签名、单据清晰度和差异；可通过、退回补件或标记争议，审核通过才进入结算。','状态：Uploaded→Reviewing→Confirmed/Returned/Disputed。','A/B [S08]'),
F('TMS-031','索赔与货损','客服、承运商、财务','基于货损/短少建立索赔，关联照片、签收差异、责任和金额；支持协商、审批、扣款或保险处理。','输出：ClaimCase、DeductionFact。','B/C'),
F('TMS-032','返程与逆向运输','客服、计划员','拒收、周转箱、退货或回单原件可生成返程需求，复用计划/委托/跟踪流程并关联原运单。','输出：ReturnTransportOrder。','B/C'),
F('TMS-033','运输计费事实','TMS、计费','按里程、重量、体积、车型、时长、节点、等待、装卸和异常生成事实；确认值与计划值分开。','输出：FreightChargeFact。','A [S04]'),
F('TMS-034','自动计费与预提','财务、系统','运输完成或节点满足后触发计费，未收到发票可按合同预提；费率缺失进入异常而不是使用默认零价。','输出：ChargeCalculation、AccrualVoucher。','A [S04]'),
F('TMS-035','承运商对账与结算','财务、承运商','按期间、承运商和合同汇总运单费用；承运商确认或争议，调整后生成应付凭证和付款状态。','输出：CarrierStatement、APVoucher。','A [S04]'),
F('TMS-036','客户运输应收','财务、客户','按客户合同和实际服务生成应收，与承运商应付独立；支持加价、最低费、燃油附加、等待费和税。','输出：CustomerStatement、ARVoucher。','A/B [S04]'),
F('TMS-037','车队与车辆运营','车队管理员','自有车维护里程、保养、维修、年检、油耗和不可用时段；不可用车辆不能进入运力池。','输出：MaintenancePlan、VehicleAvailability。','A [S04]'),
F('TMS-038','IoT 与温控监测','运输、质量','接收温度、湿度、门磁和设备告警，关联运单和时间段；超阈值触发异常，原始遥测按保留策略归档。','输出：Telemetry、ConditionAlert。','A [S04]'),
F('TMS-039','运输 KPI 与控制塔','运输主管、管理层','统计准时提货/交付、拒单率、装载率、里程、成本、异常、POD 时效和承运商绩效，可按客户/区域/线路下钻。','输出：TransportMetric。','A [S04][S06]'),
F('TMS-040','客户追踪页面','客户、收货方','通过登录或受控追踪码查看运单状态、节点、ETA、地图、异常提示和签收证明；链接过期、脱敏并限制下载。','输出：PublicTrackingView。','A [S01][S04]'),
]
FEATURE_SECTIONS.append(('TMS 运输管理','从运输需求、计划配载、委托派车，到在途跟踪、回单和自动结算。',tms_features))

ams_features = [
F('AMS-001','容量日历与资源模型','预约管理员','按仓库、月台、区域、班组或服务类型定义每日/班次容量，支持数量、托数、车数、工时等单位。','输出：CapacityCalendar、CapacityProfile。','A [S05]'),
F('AMS-002','时隙生成与关闭','系统、预约管理员','根据工作日历、班次、资源、提前期和黑名单生成时隙；可临时关闭、扩容或保留内部容量，变更记录原因。','输出：TimeSlot。','A [S05]'),
F('AMS-003','工作量计算','系统','从订单行、数量、包装、车型、装卸方式和历史效率计算预约工作量；计算规则版本化，人工修改需备注。','输出：WorkloadEstimate。','A/B [S05]'),
F('AMS-004','照单预约','客户、供应商、承运商','选择可预约订单和数量，系统校验未预约量、地点、时间窗和时隙容量；订单可分多次预约。','输出：Appointment、OrderLink。','A [S05]'),
F('AMS-005','无单预约','承运商、门岗','允许先预约车辆/服务类型，后续绑定订单；需要更严格的审批和现场识别规则，避免占用无效容量。','输出：UnlinkedAppointment。','A [S05]'),
F('AMS-006','循环/长期预约','大客户、供应商','按周几、频率和有效期创建循环预约，逐次实例占用容量；节假日或冲突实例进入待确认。','输出：RecurringAppointment、Occurrence。','A [S05]'),
F('AMS-007','可约时隙查询','预约方','按仓库、业务类型、工作量和车辆查询可用时隙；仅返回可见容量和版本，防止泄露其他客户预约。','输出：SlotAvailability。','A [S05]'),
F('AMS-008','并发容量占用','系统','提交时使用 slot version 和原子条件更新，确保 used+reserved+workload≤capacity；冲突返回新候选而非超卖。','输出：CapacityReservation。','C'),
F('AMS-009','预约提交与审批','预约方、审批人','根据客户、无单、超容量、临时加急和服务类型决定自动确认或人工审批；审批可建议改期。','状态：Draft→Submitted→Pending/Confirmed/Rejected。','A [S05]'),
F('AMS-010','预约改期','预约方、管理员','原时隙与新时隙采用一个业务事务编排：先保留新容量，再释放旧容量；失败保持原预约不变。','输出：RescheduleRecord。','A/B [S05]'),
F('AMS-011','预约取消','预约方、管理员','按提前期、状态和合同判断是否可取消及是否产生费用；释放容量并通知订单、运输和门岗。','状态：Confirmed→Cancelled。','A/B [S05]'),
F('AMS-012','到场提醒与准备','系统、仓库、司机','预约前按规则向仓库、承运商和司机发送提醒，包含地址、证件、二维码、禁限行和注意事项。','输出：ReminderSchedule、Notification。','A/B [S05]'),
F('AMS-013','门岗签到与身份校验','门岗、司机','扫描二维码/车牌或输入预约号，校验时段、车辆、司机、证件和订单；早到、迟到、无预约按策略处理。','输出：GateEvent、VerificationResult。','A [S05][S10]'),
F('AMS-014','通行证与门禁联动','门岗、设备','签到通过生成短期二维码/车牌通行证；门禁事件回写，重复入场或过期拒绝。','输出：GatePass、AccessEvent。','B/C'),
F('AMS-015','排队取号与优先级','门岗、现场调度','为待作业车辆生成排队号，优先级考虑预约时段、加急、温控和现场状态；叫号超时可顺延或取消。','输出：QueueTicket。','A [S05]'),
F('AMS-016','月台自动/人工分配','现场调度、系统','按业务类型、车型、设备、仓门、作业区和预计时长选择月台；月台冲突、故障和切换保留历史。','输出：DockAssignment。','A [S05]'),
F('AMS-017','叫号与司机通知','现场调度、司机','月台可用后通过大屏、短信、APP 或语音叫号；司机确认后进入靠台，未响应触发重叫和升级。','输出：CallEvent、DriverAcknowledgement。','A/B [S05]'),
F('AMS-018','装卸作业事件','仓库、现场调度','记录靠台、开始、暂停、恢复、完成和离台时间及数量；与 WMS/TMS 任务关联，计算等待与作业时长。','输出：OperationEvent。','A [S05]'),
F('AMS-019','出场与预约关闭','门岗、系统','校验装卸完成、单据、封签和异常后出场；释放月台，关闭预约并发布实际时长和绩效事件。','状态：Operating→CheckedOut→Completed。','A [S05]'),
F('AMS-020','迟到、爽约与处罚','系统、管理员、财务','超过宽限期未签到标记爽约；记录原因、通知和容量浪费，按合同生成罚金事实，可申诉和豁免。','输出：NoShowCase、PenaltyChargeFact。','A/B [S05]'),
F('AMS-021','现场与容量看板','仓库主管、管理层','显示今日预约、到场、排队、月台占用、作业时长、迟到、爽约和未来容量；支持按客户/承运商/业务类型下钻。','输出：AppointmentDashboard。','A [S05][S06]'),
]
FEATURE_SECTIONS.append(('AMS 预约与现场管理','容量规划、在线预约、门岗排队、月台作业和现场绩效。',ams_features))
billing_features = [
F('BIL-001','计费事实接收与去重','WMS、TMS、AMS、计费','订阅收货、存储、操作、运输、等待、爽约等事实；按 eventId/businessRef/chargeType 去重，原事实不可覆盖，只能更正。','输出：ChargeFact、FactCorrection。','A/B [S02][S03][S04]'),
F('BIL-002','合同与费率版本匹配','计费、商务','按业务发生时间、客户/供应商/承运商、组织、路线、服务和币种匹配有效合同及费率版本；多命中按优先级，零命中进入异常。','输出：RateMatch、MatchTrace。','A/B [S04]'),
F('BIL-003','阶梯、最低费与封顶','计费管理员','支持按数量、重量、体积、里程、时间和区间计费，以及最低费、封顶、起步价、进位和多条件阶梯。','输出：ChargeCalculationLine。','A/B [S04]'),
F('BIL-004','附加费与条件费用','计费管理员','支持燃油、等待、装卸、偏远、夜间、节假日、温控、保险、材料和罚金等；条件和基础金额引用可追溯。','输出：AccessorialCharge。','A/B [S04]'),
F('BIL-005','税率、币种与汇率','财务','金额使用 Money 类型；按税务场景计算含/未税，跨币种使用指定日期汇率并保存来源、舍入规则和差额。','输出：TaxDetail、FxConversion。','C'),
F('BIL-006','计费计算追踪','财务、内审','每次计算保存输入事实、匹配费率、表达式、阶梯、舍入、税和输出；费率重算生成新版本，不覆盖原结果。','输出：CalculationTrace。','C'),
F('BIL-007','应收/应付凭证','财务','把计算结果按客户/供应商/承运商、期间和业务类型生成 AR/AP 凭证；同一计算行只能进入一个有效凭证。','状态：Draft→Calculated→Validated→Approved。','A/B [S02][S04]'),
F('BIL-008','预提与冲销','财务','业务完成但发票未到时生成预提；实际账单确认后自动冲销或调整差额，预提与实际行建立引用。','输出：AccrualVoucher、ReversalVoucher。','A/B [S04]'),
F('BIL-009','凭证校验与审批','财务、审批人','校验合同、税、金额、重复、期间和业务状态；超过阈值、手工调整或无合同费用走审批。','输出：VoucherValidation、ApprovalTask。','C'),
F('BIL-010','对账单生成与发布','财务、客户/承运商','按期间和合同汇总凭证行，附业务明细和附件；发布后外部方确认或逐行提出差异。','输出：ReconciliationStatement。','A [S04]'),
F('BIL-011','对账差异与争议','财务、伙伴','差异按漏单、费率、数量、服务、税和重复分类；可接受、驳回、补证或调整，所有沟通和版本留痕。','状态：Reconciled/Disputed/Adjusted。','A/B [S04]'),
F('BIL-012','调整、扣款与分摊','财务','通过调整单增加/减少金额，支持索赔扣款、跨订单分摊和成本中心分摊；调整引用原凭证且需审批。','输出：AdjustmentVoucher、AllocationDetail。','C'),
F('BIL-013','开票/收票管理','财务','关联对账单和发票号码、税额、开票方、日期与附件；支持部分开票、多票和红冲，避免超额。','输出：Invoice、InvoiceLine。','C'),
F('BIL-014','收付款与核销','财务、ERP','从财务系统同步或录入收付款，按发票/凭证核销；支持部分付款、退款、手续费和未达账。','输出：Payment、SettlementAllocation。','C'),
F('BIL-015','期间关闭与重开','财务主管','关闭期间前检查未审批、争议、未开票和汇率；关闭后禁止普通修改，重开需高权限和审计。','状态：Open→Closing→Closed/Reopened。','C'),
F('BIL-016','结算报表与毛利','财务、管理层','按客户、承运商、仓库、线路、订单和服务比较收入、成本、预提、实际和毛利；指标可追溯到凭证和计算行。','输出：MarginMetric、SettlementReport。','A/B [S04][S06]'),
]
FEATURE_SECTIONS.append(('计费与结算中心','业务事实到费率计算、应收应付、对账、发票和收付款。',billing_features))

control_features = [
F('CTL-001','统一业务事件模型','平台、各领域','所有事件包含 eventId、tenantId、aggregateType/id/version、eventType、occurredAt、traceId、schemaVersion 和 payload。','输出：BusinessEvent。','C'),
F('CTL-002','端到端业务时间线','客服、运营、管理层','按订单号聚合 OMS、WMS、TMS、AMS、结算和集成事件，支持因果关系、来源、重放状态和附件下钻。','输出：EndToEndTimeline。','A/B [S06]'),
F('CTL-003','订单可视化','客服、管理层','展示订单确认、分配、仓储、运输、签收和结算的当前状态、完成率、承诺时间和阻塞点。','输出：OrderControlView。','A/B [S02][S06]'),
F('CTL-004','库存网络可视化','库存计划、管理层','按仓库、货主、商品、区域显示可用、冻结、在途、库龄、周转和缺货风险；数据标注刷新时间。','输出：InventoryNetworkView。','A/B [S06]'),
F('CTL-005','运输网络与地图','运输、管理层','地图展示运单、车辆、线路、节点、ETA、延误和温控异常；支持聚合、热力和钻取。','输出：TransportNetworkView。','A [S04][S06]'),
F('CTL-006','预约与月台可视化','仓库、管理层','展示未来容量、今日到场、排队、月台占用、作业时长、迟到和爽约，联动 WMS/TMS 状态。','输出：YardControlView。','A/B [S05][S06]'),
F('CTL-007','SLA 计时器','系统、运营','为确认、释放、收货、上架、发运、交付、回单和对账建立计时器；日历、暂停、重开和责任域可配置。','输出：SlaClock、SlaEvent。','B/C'),
F('CTL-008','预警规则与复杂条件','运营管理员','基于事件、状态、持续时间、指标和业务属性定义条件；支持去抖、抑制、合并和租户/客户覆盖。','输出：AlertRule、RuleVersion。','A [S06]'),
F('CTL-009','例外工单与责任路由','运营、客服','预警生成例外工单，根据业务域、组织、客户、严重度和轮班路由；保存责任人、截止时间和状态。','输出：AlertCase、AssignmentHistory。','A/B [S06]'),
F('CTL-010','告警通知与升级','系统、管理层','按严重度和 SLA 通过站内、邮件、短信、微信或推送通知；未确认/未解决自动升级到主管。','输出：Notification、EscalationEvent。','A [S06]'),
F('CTL-011','根因与处置知识库','运营、质量','关闭异常时选择根因、责任和解决方案；相似案例可检索，形成规则优化、培训和供应商改进建议。','输出：RootCauseRecord、KnowledgeArticle。','B/C'),
F('CTL-012','KPI 指标定义','BI 管理员','指标定义包含公式、粒度、维度、时区、数据源和刷新策略；版本变更不重写历史口径，报表显示口径版本。','输出：MetricDefinition。','A/B [S06]'),
F('CTL-013','运营看板与大屏','管理层、运营','可配置卡片、趋势、地图、排行、漏斗、异常和滚动信息；支持多组织/仓库切换和大屏自动刷新。','输出：Dashboard、WidgetConfig。','A [S06]'),
F('CTL-014','自助分析与报表','分析师','在受控语义模型上选择指标、维度和筛选，保存/分享报表并导出；复杂查询受配额和脱敏约束。','输出：ReportDefinition、QueryJob。','A/B [S06]'),
F('CTL-015','数据湖与历史快照','数据团队','通过 CDC 和事件构建事实/维度模型，保留业务状态快照和迟到数据处理；重算与在线事务隔离。','输出：FactTable、DimensionSnapshot。','C'),
F('AI-001','车辆路径优化','计划员','在时间窗、容量、司机工时、路况和成本约束下优化车辆数、里程或准时率；结果保存约束和解释。','输出：RouteOptimizationResult。','A [S07]'),
F('AI-002','装载优化','运输、仓库','根据箱体/托盘尺寸、重量、重心、禁配和装卸顺序生成装载建议；建议需人工确认并可反馈实际偏差。','输出：LoadOptimizationResult。','A [S07]'),
F('AI-003','需求预测','计划、管理层','按商品、客户、区域和时间预测订单/出库需求，记录训练数据窗口、版本、置信区间和偏差。','输出：DemandForecast。','A [S07]'),
F('AI-004','库存与补货优化','库存计划','结合需求、服务水平、前置期、成本和保质期建议安全库存、补货点和调拨；建议不直接改库存。','输出：InventoryPolicyRecommendation。','A [S07]'),
F('AI-005','供应链网络与情景模拟','管理层','模拟仓库开关、承运能力、需求变化、成本和 SLA 对网络的影响；情景与基线独立，可比较并导出。','输出：NetworkScenario、SimulationResult。','A [S07]'),
]
FEATURE_SECTIONS.append(('控制塔、BI 与优化','端到端可视化、事件预警、例外治理、指标、大屏和 AI 优化。',control_features))

integration_features = [
F('INT-001','API Gateway 与流量治理','集成、运维','统一 TLS、鉴权、租户解析、限流、IP 策略、请求大小和 correlationId；敏感接口可配置更严格配额。','输出：GatewayLog、RateLimitDecision。','C'),
F('INT-002','OAuth/API Key/签名认证','集成管理员','支持 OAuth2 客户端、API Key、mTLS 和 HMAC 签名；凭证可轮换、限定作用域、IP 和有效期。','输出：ApiCredential、AccessToken。','C'),
F('INT-003','OpenAPI 与版本管理','开发者、伙伴','发布可机器读取的接口契约、示例、错误码和弃用日期；新增字段向后兼容，破坏性变更使用新版本。','输出：ApiDefinition、DeprecationNotice。','A/B [S02]'),
F('INT-004','写接口幂等','外部系统','外部创建/确认/取消请求携带 Idempotency-Key；相同键相同内容返回原结果，不同内容返回冲突。','输出：IdempotencyRecord。','C'),
F('INT-005','EDI 与 SFTP 文件交换','集成、伙伴','支持约定格式、目录、命名、加密、回执和归档；文件级与行级状态分离，重复文件通过摘要去重。','输出：FileExchange、AckMessage。','A/B [S02][S04]'),
F('INT-006','Webhook 订阅','伙伴、开发者','伙伴按事件类型和对象范围订阅；签名、重试、退避、禁用和死信可配置，投递至少一次。','输出：WebhookSubscription、DeliveryAttempt。','B/C'),
F('INT-007','字段映射与转换','实施顾问','通过版本化映射定义外部字段、字典、单位、地址和默认值；映射发布前用样例测试，运行记录版本。','输出：MappingDefinition、TransformResult。','B/C'),
F('INT-008','集成消息监控与重放','集成运维','查看入站/出站消息、业务键、状态、耗时、错误和重试；修复映射/主数据后可授权重放，保留原消息。','输出：IntegrationMessage、ReplayRecord。','B/C'),
F('INT-009','ERP/财务系统适配器','集成','同步主数据、订单、库存、凭证、发票和收付款；适配器隔离不同厂商语义，不让 ERP 字段渗透领域模型。','输出：AdapterCommand/Event。','A/B [S02][S03][S04]'),
F('INT-010','设备与 IoT 网关','设备、运维','接入 RF、打印、称重、门禁、GPS 和温控；设备注册、心跳、证书、命令回执和遥测限流统一管理。','输出：Device、Telemetry、DeviceCommand。','A/B [S03][S04]'),
F('MOB-001','客户移动端','客户','查看订单、库存、预约、运输、签收和对账；支持确认、取消、退货、消息和附件，权限与 Web 一致。','输出：CustomerMobileView。','A [S01]'),
F('MOB-002','司机移动端','司机','任务列表、导航、节点上报、扫码、照片、异常、电子签名和回单；离线可用并限制设备绑定。','输出：DriverTask、OfflineQueue。','A [S01][S04]'),
F('MOB-003','仓库 RF/PDA','仓库作业员','任务领取、扫描、数量确认、异常、打印和离线同步；操作界面强调大按钮、少输入和声光反馈。','输出：WarehouseTaskMobileView。','A [S01][S03]'),
F('MOB-004','合作伙伴门户','供应商、承运商','供应商处理订单/ASN/预约，承运商处理委托、车辆司机、跟踪、回单和对账；数据按伙伴范围隔离。','输出：PartnerPortalView。','A/B [S02][S04]'),
F('OPS-001','结构化日志、指标与链路','运维、开发','日志包含 tenantId、userId、businessRef、traceId；关键命令、事件、外部依赖和数据库调用产生指标与 trace。','输出：Log、Metric、Trace。','C'),
F('OPS-002','业务与技术监控','运维、业务运营','监控 API 错误、队列积压、任务延迟、数据库、缓存、设备和第三方服务，同时监控未释放订单、任务积压和 SLA。','输出：MonitorAlert。','C'),
F('OPS-003','备份、恢复与容灾','运维、安全','数据库时间点恢复、对象存储版本、跨区备份和定期恢复演练；明确 RPO/RTO，灾备切换不产生双写。','输出：BackupSet、DRDrillReport。','C'),
F('OPS-004','CI/CD 与数据库迁移','开发、运维','自动测试、镜像扫描、灰度发布和回滚；数据库迁移采用向前兼容的 expand/migrate/contract，不在高峰锁大表。','输出：Release、MigrationRun。','C'),
F('OPS-005','数据保留与归档','安全、法务、运维','按业务、审计、轨迹、附件和遥测定义保留期；归档可检索，删除遵守法律留置和租户合同。','输出：RetentionPolicy、ArchiveJob。','C'),
F('OPS-006','隐私、脱敏与主体请求','安全、法务','识别个人信息字段，按角色脱敏和限制导出；支持访问/更正/删除请求，但交易、审计和法定义务数据按规则保留。','输出：PrivacyRequest、RedactionRecord。','C'),
F('OPS-007','容量、性能与成本治理','运维、架构师','按租户和业务域统计请求、存储、事件、报表和导出资源；设置配额、冷热分层和扩缩容策略。','输出：CapacityPlan、UsageMetric。','C'),
F('OPS-008','租户配置与数据迁移工具','实施、运维','支持从试点到生产迁移配置、主数据和未完业务；迁移前校验依赖，迁移后做数量/金额/状态对账。','输出：MigrationPlan、ReconciliationReport。','C'),
]
FEATURE_SECTIONS.append(('开放集成、移动端与运维','开放接口、伙伴/移动协同、设备接入和生产级可运营性。',integration_features))

ROLE_ROWS = [
['平台管理员','租户开通、套餐、全局配置、平台运维','不可查看租户业务明细，除非使用受审计的支持会话。'],
['租户管理员','组织、用户、角色、数据范围、租户配置','不能绕过业务审批或直接修改业务流水。'],
['主数据管理员','商品、伙伴、地址、仓库、合同、费率','关键变更需审批，不能修改历史业务快照。'],
['订单客服/运营','订单录入、审核、冻结、分配、变更、异常','只能在授权客户/组织范围内操作。'],
['仓库主管','任务、波次、库存、盘点、人员和设备','库存调整、解冻和发运确认需更高动作权限。'],
['仓库作业员','RF 收货、上架、拣选、包装、装车','只看到已分配或可领取任务，关键异常需主管确认。'],
['运输计划员/调度员','运输计划、配载、承运商、车辆司机、改派','费用超阈值、非合同承运商和超载方案需审批。'],
['承运商管理员','接单、报价、车辆司机、跟踪、回单、对账','只能访问分配给本伙伴的数据。'],
['司机','本人任务、导航、节点、异常、签收、POD','不能查看商业费率和其他司机任务。'],
['预约管理员/门岗','容量、时隙、审批、签到、排队、月台','容量扩容、爽约豁免需额外权限。'],
['财务','计费、凭证、对账、发票、收付款、期间关闭','费率维护与凭证审批应职责分离。'],
['控制塔运营','事件、SLA、告警、例外工单、看板','补偿动作仍需调用各业务域命令并鉴权。'],
['客户/供应商','本方订单、ASN、预约、追踪、签收、对账','数据范围绑定伙伴 ID，内部评论不可见。'],
['审计/只读','审计日志、报表和历史快照','无业务写权限，导出受控并留痕。'],
]

UI_ROUTE_ROWS = [
['工作台','/workbench','待办、快捷入口、最近访问、收藏、KPI 摘要','工作台模板'],
['租户/组织/用户/角色','/admin/*','租户、组织树、用户、角色、权限、数据范围','列表+树+抽屉'],
['配置/工作流/规则','/admin/config|workflows|rules','配置版本、流程设计、规则模拟与发布','设计器+版本历史'],
['主数据','/mdm/products|partners|warehouses|contracts','商品、伙伴、地址、仓库、合同、费率','主从编辑'],
['OMS 订单','/oms/orders','订单查询、详情、审核、冻结、分配、释放','列表+详情'],
['OMS 分配/异常','/oms/allocation|exceptions','分配决策、欠货、例外工作台','工作台'],
['WMS 入库','/wms/inbounds|receiving|inspection|putaway','预期、收货、质检、上架','列表+任务'],
['WMS 库存','/wms/inventory|counts|transfers','库存、批次序列、盘点、移库、调整','查询+任务'],
['WMS 出库','/wms/outbounds|waves|picks|packing|loading','出库、波次、分配、拣选、包装、装车','列表+任务'],
['TMS 运输','/tms/orders|planning|shipments|tracking','订单池、计划、委托、派车、跟踪','计划板+列表+地图'],
['AMS 预约','/ams/calendar|appointments|yard','容量、预约、门岗、排队、月台','日历+现场看板'],
['结算中心','/billing/calculations|vouchers|reconciliation','计费、凭证、对账、发票、收付款','列表+明细'],
['控制塔','/control/orders|inventory|transport|alerts','端到端视图、地图、告警、SLA、例外','大屏+工作台'],
['集成中心','/integration/endpoints|messages|jobs','端点、映射、消息、重试、回放','监控台'],
['报表中心','/reports|dashboards','保存报表、指标、订阅、导出','报表设计器'],
]

COMPONENT_ROWS = [
['ApplicationShell','左侧深色模块导航；顶部租户/组织/仓库/语言/用户；多标签内容区；全局消息和命令面板。'],
['QueryPanel','字段化筛选、快速筛选、日期范围、展开/收起、重置、保存视图；查询条件序列化到 URL 或页面会话。'],
['CommandBar','新建、审核、冻结、释放、导入、导出、打印等动作；根据资源、状态和数据权限动态启用。'],
['DataGrid','服务端分页/排序/筛选；固定列、列显隐、聚合、批量选择、行上下文菜单、导出当前视图。'],
['MasterDetail','左侧列表或主表，右侧抽屉/下方子表展示明细；保存时带 version，冲突提示比较。'],
['StatusBadge/Stepper','统一颜色与状态词；状态阶梯显示已完成、当前、不可达和异常节点。'],
['Timeline','跨域事件按时间、来源、操作者和 traceId 展示；支持过滤、附件和原始消息权限查看。'],
['TaskWorkbench','待处理队列、领取/转派、扫描步骤、异常、计时和绩效；支持键盘/扫码枪快速操作。'],
['Map/Gantt/Calendar','运输地图、计划甘特、预约容量日历使用同一选择和详情抽屉约定。'],
['Drawer/Modal','抽屉用于上下文编辑，模态框用于短决策；复杂业务不在多层模态中完成。'],
['Toast/NotificationCenter','短反馈使用 Toast；可行动或需追踪的消息进入消息中心/例外工单。'],
['FieldUpdater','批量更新所选对象的允许字段，展示影响数量、权限和逐行失败，不支持绕过状态机。'],
]

FLOW_DETAILS = [
('订单到交付与结算','10_e2e_order',['客户/ERP','OMS','WMS','TMS','AMS','结算','控制塔'],'主数据、合同、仓库、承运能力已配置。','订单状态由 OMS 拥有；WMS/TMS/AMS 通过事件报告事实；结算只消费已确认事实。','校验失败、库存不足、任务短缺、承运商拒单、延误和签收差异进入统一例外工单。'),
('WMS 入库','11_inbound',['供应商/采购','AMS/门岗','收货员','质检员','上架员'],'入库单或 ASN 已发布。','收货、质检和上架分别建任务；库存只有在确认动作中变更；批次/序列规则强校验。','超短收、破损、效期不符、质检不合格和上架容量不足。'),
('WMS 出库','12_outbound',['OMS','波次计划员','拣选/包装/装车','TMS'],'出库单已释放且库存可分配。','分配不能超过可用库存；短拣必须原因码；发运确认是不可逆业务边界。','缺货、错拣、称重差异、漏装、车辆/月台不匹配。'),
('TMS 运输','13_tms',['运输计划员','承运商','调度/司机','客户','控制塔'],'运输需求和伙伴/车辆资料可用。','计划、审批、委托、接受、派车、发运和签收为独立状态；轨迹事件至少一次投递。','拒单/超时、证照到期、偏航、延误、温控、拒收和 POD 退回。'),
('AMS 预约','14_appointment',['预约方','AMS','审批人','门岗','现场调度','仓库'],'容量日历、资源和工作量规则已发布。','时隙占用必须原子校验版本；改期先占新再释旧；现场事件与预约状态一致。','容量冲突、无单审批、迟到、爽约、证件失败、月台故障。'),
('库存盘点与移库','15_inventory_transfer',['库存主管','盘点/移库员','审批人'],'库存余额和任务无冲突。','库存变更只写流水并派生余额；盘点差异审批后入账；并发操作使用版本。','冻结失败、序列号冲突、目标容量不足、账实差异。'),
('退货与逆向','16_return',['客户','OMS','TMS/AMS','WMS','财务'],'原订单存在且在退货政策内。','RMA 约束可退数量；退货接收后按质检处置；退款与库存动作解耦但可追溯。','无授权、超期、货不符、质检失败、退款争议。'),
('越库','17_crossdock',['OMS/WMS','收货员','出库/装车','TMS'],'存在时间匹配的入库供应和出库需求。','越库匹配必须满足商品、批次、数量、质检和时间窗；失败回退普通入库。','短收、晚到、质检不通过、运输班次变化。'),
('例外与补偿','18_exception',['各领域','控制塔','责任人','管理层'],'事件或监控识别异常。','例外工单不直接改业务表；所有修复通过领域命令，关闭前验证状态与数据一致性。','重复告警、责任不清、补偿失败、超 SLA 升级。'),
('计费与结算','19_settlement',['业务域','计费','财务','客户/承运商','ERP'],'计费事实、合同和费率版本可用。','费率按发生时点匹配；计算不可覆盖；调整引用原凭证；期间关闭后受控重开。','费率缺失、重复事实、金额异常、对账争议、发票差异。'),
]

SEQUENCE_DETAILS = [
('外部订单接入与跨域编排','20_order_api_seq','同步部分只完成鉴权、校验、去重和 OMS 事务；WMS/TMS/控制塔通过事件异步接收。API 返回 202 并允许按 orderId 查询最终状态。'),
('RF 扫描与库存并发','21_wms_rf_seq','RF 命令携带 taskVersion、设备序号和幂等键；库存聚合通过原子更新或乐观锁保证不超扣，离线重放发生冲突时不静默覆盖。'),
('司机执行与 POD','22_driver_seq','司机节点、GPS、ETA、客户门户和控制塔均围绕同一运单事件流；附件先上传对象存储，再提交 POD 元数据。'),
('预约容量与现场执行','23_appointment_seq','可用时隙只是读快照，最终提交必须再次原子占用；改期/取消向 WMS、TMS、门岗发布事件。'),
('计费、对账与财务','24_billing_seq','计费服务消费不可变事实，匹配生效时点费率，保存计算追踪；凭证、对账和发票分别有状态机。'),
]

STATE_RULE_ROWS = [
['订单','Draft→Open→Approved→Allocated→Released→Executing→Completed→Closed','数量守恒；已执行部分不可直接取消；Completed 需要所有履约行终态。'],
['入库','Draft→Expected→Arrived→Receiving→Received→Inspecting/Putaway→Completed','已收数量=接受+拒收+待定；可用库存只来自放行且已上架数量。'],
['出库','Draft→Released→Waved→Allocated→Picking→Packed→Staged→Loaded→Shipped','分配≤可用；已发运不可逆；短拣必须释放或重分配差额。'],
['运输','Open→Planned→Approved→Tendered→Accepted→Dispatched→Tracking→Delivered→POD→Settled→Closed','承运商接受后改派需显式撤销；节点时间单调；结算要求 POD/例外处理满足策略。'],
['预约','Draft→Submitted/Pending→Confirmed→CheckedIn→Queued/Docked→Operating→Completed','容量不超售；一个预约同一时刻最多一个有效月台；改期不丢失原容量。'],
['凭证','Draft→Calculated→Validated→Approved→Reconciled→Invoiced→Paid→Closed','计算行不能重复入账；关闭期间禁止普通修改；调整通过反向/新凭证。'],
]

RULE_ROWS = [
['订单分配','订单行、库存、仓库、服务区、产能、成本','过滤不可用候选→按硬约束校验→按目标函数评分→分配并预占','保存候选、排除原因、规则版本；失败产生 Backorder/Exception。'],
['上架策略','商品、LPN、批次、库位容量、温层、混放规则','过滤禁用/不兼容库位→同品同批/固定优先→距离与补货评分','保存推荐与人工改选；改选仍需硬约束校验。'],
['波次/分配','出库截止、路线、客户、商品、库存、产能','选择订单→模拟工作量→分配库存→生成任务','分配原子更新；不足转缺货流程。'],
['承运商选择','线路、车型、服务、合同、价格、绩效、容量','过滤资质/容量→匹配合同→评分或竞价→审批/委托','拒单释放容量；保存报价和选中原因。'],
['预约时隙','工作量、资源容量、提前期、客户规则','生成候选→读可用量→提交时原子占用→审批/确认','version 冲突返回新候选；不能先释放旧时隙再改期。'],
['计费','事实、合同、费率版本、税、币种','按发生时点匹配→计算基础费/阶梯/附加→税和舍入→校验','完整 calculationTrace；零命中/多命中进入异常。'],
]

AGGREGATE_ROWS = [
['BusinessOrder','OrderLine、Address、Party、Hold','订单行数量与订单状态一致；正式订单变更走 Change；状态只通过命令。','OMS 拥有；发布 order.* 事件。'],
['InboundOrder','InboundLine、ReceiptTask 引用','收货数量分解守恒；完成要求所有行已上架/处置。','WMS 入库拥有；发布 inbound.*。'],
['InventoryBalance','余额键、Reservation、Movement','available=onHand-allocated-hold；version 单调；流水不可变。','WMS 库存拥有；发布 inventory.*。'],
['OutboundOrder/Wave','OutboundLine、Allocation、Pick/Pack/Load 引用','分配不超可用；发运后不可逆；任务完成量汇总到订单行。','WMS 出库拥有；发布 outbound.*。'],
['Shipment','ShipmentItem、Stop、Tender、Assignment、POD','至少两站；节点顺序和实际时间单调；一个有效车辆指派。','TMS 拥有；发布 shipment/tracking.*。'],
['Appointment','Slot 引用、OrderLink、Queue、DockAssignment','容量不超售；一个有效时隙和月台；终态释放资源。','AMS 拥有；发布 appointment.*。'],
['SettlementVoucher','VoucherLine、Reconciliation/Invoice 引用','计算行不重复入账；金额=行合计；关闭后只通过调整。','Billing 拥有；发布 voucher.*。'],
['AlertCase','AlertRule 引用、Action、Notification','同一去重键在窗口内合并；关闭需根因和解决验证。','Control Tower 拥有。'],
]

API_ROWS = [
['POST','/api/v1/oms/orders','创建订单','Idempotency-Key；返回 orderId/version/accepted。'],
['POST','/api/v1/oms/orders/{id}/approve','审核订单','If-Match/version；状态与权限校验。'],
['POST','/api/v1/oms/orders/{id}/release','释放履约','可同步受理，跨域任务异步。'],
['GET','/api/v1/oms/orders','订单查询','服务端分页、稳定排序、保存视图参数。'],
['POST','/api/v1/wms/inbounds/{id}/receive','提交收货','taskVersion、扫描明细、设备序号。'],
['POST','/api/v1/wms/tasks/{id}/confirm','确认仓库任务','幂等、乐观锁、库存原子更新。'],
['GET','/api/v1/wms/inventory','库存查询','返回 snapshotAt；承诺需另走预占命令。'],
['POST','/api/v1/tms/plans/{id}/approve','批准运输计划','校验负载、成本、资质和版本。'],
['POST','/api/v1/tms/shipments/{id}/events','上报运输节点','事件时间、位置、来源、幂等键。'],
['POST','/api/v1/tms/shipments/{id}/pod','提交 POD 元数据','附件对象引用；异步病毒扫描/审核。'],
['GET','/api/v1/ams/slots','查询可约时隙','读快照和 slotVersion。'],
['POST','/api/v1/ams/appointments','创建预约','提交时原子容量占用。'],
['POST','/api/v1/billing/calculations','触发重算','引用事实和费率版本；生成新计算版本。'],
['GET','/api/v1/control/timelines/{businessRef}','端到端时间线','跨域查询投影，不跨库事务。'],
]

EVENT_ROWS = [
['order.created.v1','OMS','orderId、channel、customer、lines、requestedWindow','控制塔、集成、规则'],
['order.released.v1','OMS','orderId、fulfillmentRequests、shipmentRequests','WMS、TMS、控制塔'],
['inbound.completed.v1','WMS','inboundId、received/accepted/rejected、inventoryRefs','OMS、ERP、计费、控制塔'],
['inventory.changed.v1','WMS','balanceKey、delta、movementType、businessRef','OMS 查询投影、控制塔、数据湖'],
['outbound.ready.v1','WMS','outboundId、packages、weight/volume、dock','TMS、控制塔'],
['outbound.shipped.v1','WMS','outboundId、shipmentId、actualAt、packages','OMS、ERP、计费'],
['shipment.tendered.v1','TMS','shipmentId、carrier、expiresAt、price','承运商门户、通知、控制塔'],
['tracking.event.v1','TMS','shipmentId、stopId、type、time、location、source','客户门户、控制塔、ETA'],
['shipment.delivered.v1','TMS','shipmentId、deliveredAt、PODRef、variance','OMS、计费、客户、控制塔'],
['appointment.confirmed.v1','AMS','appointmentId、slot、vehicle、orderLinks','WMS、TMS、门岗、通知'],
['appointment.completed.v1','AMS','appointmentId、actualTimes、workload、noShow/overrun','WMS/TMS、计费、控制塔'],
['charge.calculated.v1','Billing','calculationId、businessRef、amount/tax、rateVersion','结算凭证、控制塔'],
['alert.opened.v1','Control','alertCaseId、businessRef、severity、owner、dueAt','通知、工作台、数据湖'],
]

NFR_ROWS = [
['可用性','核心下单/仓储/运输执行建议月度 99.9% 或按合同；报表和优化可较低等级。','多可用区、无状态扩展、健康检查、依赖熔断、降级读模型。'],
['性能','常用列表 P95 < 2s；普通命令 P95 < 1.5s（不含外部依赖）；RF 确认 P95 < 800ms。','服务端分页、覆盖索引、读写分离、缓存、异步长任务。'],
['吞吐','按业务峰值分别估算订单/扫描/轨迹/事件；轨迹与遥测必须独立伸缩。','事件分区、批处理、背压、冷热存储。'],
['一致性','域内强一致；跨域最终一致，关键编排有 Saga/补偿和可见状态。','事务 Outbox、消费者 Inbox、幂等键、版本和对账任务。'],
['安全','最小权限、MFA、加密、审计、密钥轮换、漏洞管理。','RBAC+ABAC、TLS/KMS、WAF、依赖扫描、渗透测试。'],
['多租户','任何请求必须有可信 tenant context；防越权作为发布门禁。','行级/Schema/独立库混合；自动化跨租户测试。'],
['可恢复性','明确 RPO/RTO；定期恢复演练和事件重放演练。','PITR、对象版本、跨区备份、Runbook。'],
['可观测','技术 trace 与订单/任务/运单/预约/凭证业务号关联。','OpenTelemetry、结构化日志、业务指标和告警。'],
['可配置性','规则、流程、费率、模板、字典发布版本不可变。','草稿/发布/回滚；业务引用版本。'],
['可访问性','键盘导航、焦点、对比度、表格语义、错误提示和移动端可读。','组件库基线和自动/人工审计。'],
]

PHASE_ROWS = [
['0. 现网核验（2-4 周）','只读账号/录屏/HAR、菜单和角色矩阵；建立页面/字段/动作清单与差异台账。','可核验需求基线、UI 样式令牌、接口依赖清单。'],
['1. 平台骨架（4-6 周）','租户、组织、用户、权限、工作台、配置、审计、附件、通知、集成框架。','可登录、多租户隔离、模块路由和统一表格/表单。'],
['2. OMS + WMS 基础闭环（8-12 周）','订单接入/审核/分配/释放，入库、库存、出库、RF，基础事件和控制塔时间线。','订单到仓储发运可演示，库存一致性通过并发测试。'],
['3. TMS + AMS（8-12 周）','运输计划/委托/司机/跟踪/POD，预约容量/门岗/月台。','订单到签收和现场预约闭环。'],
['4. 结算与控制塔（6-10 周）','费率、计费、凭证、对账、预警、KPI、大屏。','业务事实到财务和管理可视化闭环。'],
['5. 优化、行业扩展与规模化（持续）','路径/装载/预测、自动化设备、行业规则、性能/容灾、租户迁移。','按真实数据迭代模型和运营指标。'],
]

VERIFY_ROWS = [
['菜单与角色','每个角色可见菜单、路由、按钮、字段和数据范围是否一致；同一页面是否因组织/租户配置变化。'],
['列表页','查询条件、默认排序、分页、列、右键菜单、底部工具栏、批量动作、保存视图和导出。'],
['详情页','所有字段、必填/只读条件、子表、状态、时间线、附件、评论、打印和快捷键。'],
['状态机','每个动作前置条件、成功状态、失败错误、撤销/回退、终态和跨域回写。'],
['接口','浏览器网络请求、URL/方法、请求/响应字段、错误码、分页、幂等、上传和轮询/推送。'],
['权限','菜单权限之外的 API、按钮、字段、导出和数据范围；越权请求必须由后端拒绝。'],
['并发','两人同时审核、分配、扫描、占时隙、指派车辆和修改费率时的冲突行为。'],
['异常','断网、重复提交、超时、第三方失败、设备离线、消息重复/乱序、任务部分成功。'],
['非功能','峰值数据量、响应时间、浏览器/设备、打印机、条码、文件大小、保留期和容灾。'],
['视觉','信息密度、栅格、字号、颜色、状态、表格、抽屉、弹窗和响应式；品牌、图标、文案与代码使用原创资产。'],
]

SOURCES = [
['S01','目标工作台入口/登录页','https://app.360scm.com/SCM.Cloud.Web/WorkStation/Index','公开可访问部分只到登录页；可确认 Web/移动端入口线索，不能确认登录后完整菜单。'],
['S02','科箭 OMS 产品页','https://www.qt-asia.com/oms.html','订单全生命周期、多渠道、分配、协同、WMS/TMS 集成、结算、开放 API 与 BI。'],
['S03','科箭 WMS 产品页','https://www.qt-asia.com/wms.html','入库、库存、出库、条码、波次、路径、增值、劳务及与 TMS/预约联动。'],
['S04','科箭 TMS 产品页','https://www.qt-asia.com/tms.html','运输订单、计划、配载、委托、跟踪、计费、移动端、预约、车队/IoT 和 KPI。'],
['S05','科箭 AMS 产品页','https://www.qt-asia.com/ams.html','容量、照单/无单/循环预约、司机移动协同、门岗排队和月台作业。'],
['S06','科箭供应链控制塔产品页','https://www.qt-asia.com/SCT.html','端到端视图、事件预警、BI/大屏、例外和优化。'],
['S07','科箭供应链 AI 产品页','https://www.qt-asia.com/AI.html','路径、装载、需求、库存和网络优化。'],
['S08','公开可访问的 Power TMS 操作手册（35 页）','https://mkp-res.hc-cdn.com/marketplace/public/app/attachment/20230531/9840555e-4526-422f-9898-5536ab0faba6.pdf','用于确认发货订单、审核、配载、运单、委托、发货、在途节点、签收与回单的历史操作链；未复制其截图或大段文字。'],
['S09','公开 WMS Cloud Overview 演示材料','https://www.slideshare.net/slideshow/power-wms-cloud-overview-cn/81367342','用于补充仓储功能全景；属于历史/第三方公开材料，证据等级为 B。'],
['S10','科箭 WMS/OMS 创新盘点文章','https://www.qt-asia.com/article/271.html','用于交叉核对预约、越库及 OMS-WMS-TMS 协同。'],
]


def build_doc(diagrams: Dict[str,Path]) -> Path:
    doc=Document(); configure_doc(doc)
    props=doc.core_properties
    props.author='OpenAI'; props.last_modified_by='OpenAI'
    props.title='SCM Cloud 同类系统 - 架构与功能设计说明书'
    props.subject='clean-room 功能逆向抽象、系统架构、工作流、时序、状态机与类图'
    props.keywords='SCM, OMS, WMS, TMS, AMS, 系统架构, 流程图, 类图'

    # Cover
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(70)
    set_run_font(p.add_run('SCM Cloud 同类系统'),15,True,'2A6FDB')
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(12)
    set_run_font(p.add_run('功能逆向抽象、系统架构与详细设计说明书'),25,True,'16324F')
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(12)
    set_run_font(p.add_run('基于公开资料的 clean-room 功能等价设计'),10,False,'5B6472',True)
    doc.add_paragraph().paragraph_format.space_after=Pt(16)
    add_table(doc,['项目','内容'],[
      ['目标入口','app.360scm.com / WorkStation（登录保护）'],
      ['文档版本','v0.1 - 公开资料基线'],
      ['生成日期','2026-07-13'],
      ['交付范围',f'{sum(len(x[2]) for x in FEATURE_SECTIONS)} 项详细功能；{len(diagrams)} 张结构/流程/时序/状态/类图；API、事件、非功能和实施计划'],
    ],widths=[4.5,12.2],font_size=8.8,header_fill='DCE9FF')
    add_callout(doc,'范围与准确性声明','目标站点当前公开入口跳转登录页，因此本文不是对供应商私有源代码、数据库或现网全部配置的复制。本文把公开产品资料、公开操作手册所能确认的能力与供应链领域建模结合，形成可独立实现的功能等价蓝图。所有推断均标注证据等级。','warn')
    add_callout(doc,'知识产权与实现方式','实现时可复现信息架构、交互效率和业务行为，但应使用原创品牌、图标、文案、样式令牌和代码；不要复制受保护的源代码、素材或绕过访问控制。','info')
    doc.add_page_break()

    # 0
    add_section_title(doc,'0','文档使用说明','先确认边界，再按“架构 - 流程 - 类图 - 功能规格 - 接口/事件 - 验收”推进实施。')
    add_table(doc,['版本','日期','说明'],[['v0.1','2026-07-13','基于公开资料、公开历史操作手册和领域建模形成；登录后细节待核验。']],widths=[2.2,3.4,11.2],font_size=8.4)
    doc.add_heading('0.1 证据等级',level=2)
    add_table(doc,['等级','含义','使用方式'],[
      ['A - 已确认','来自官方公开产品页或公开操作手册明确描述。','可作为能力范围基线，但具体字段/规则仍需登录后核验。'],
      ['B - 强推断','由公开界面/手册、多个官方页面或典型产品行为交叉推断。','可进入原型和技术设计，验收前必须核对。'],
      ['C - 建议实现','为保证一致性、安全、可扩展和可运营性而设计的 clean-room 方案。','不代表目标系统现有实现；作为自建系统的工程基线。'],
    ],widths=[2.8,7.0,7.0],font_size=8.2)
    doc.add_heading('0.2 文档导航',level=2)
    add_bullets(doc,['1 执行摘要与边界','2 能力地图与系统上下文','3 建议系统架构','4 前端信息架构与页面规范','5 角色、权限和授权模型','6 详细功能规格','7 业务流程图','8 关键时序图','9 状态机与业务规则','10 领域模型与类图','11 API、事件与一致性','12 数据、安全与非功能','13 实施路线与验收','14 现网核验清单','15 参考资料'])

    # 1
    add_section_title(doc,'1','执行摘要与边界')
    add_body(doc,'公开资料能够确认该产品族覆盖 OMS、WMS、TMS、AMS、结算、供应链控制塔、移动应用、开放集成、IoT 与优化能力。目标工作台当前只公开登录页，现网菜单、字段、角色差异、客户定制和最新视觉样式无法在未授权状态下完整观察。[S01]-[S10]')
    add_body(doc,'建议采用“逻辑服务化、物理渐进”的路线：领域边界、事件契约和数据所有权按服务设计；首期可用模块化单体降低分布式复杂度，达到组织和负载阈值后再拆分。')
    add_callout(doc,'核心难点','页面数量不是主要风险。真正困难的是订单承诺与执行、库存并发与离线扫描、运输计划与实时事件、预约容量与现场资源、费率版本与财务可追溯这五类一致性。本文所有状态机、聚合和事件契约都围绕这些不变量设计。','success')
    doc.add_heading('1.1 已确认的能力边界',level=2)
    add_bullets(doc,[
      'OMS：多渠道订单、全生命周期、分配/拆合、协同、WMS/TMS 集成、结算和可视化。[S02]',
      'WMS：入库、库存、出库、条码、波次/拣选路径、增值服务、劳务和设备集成。[S03][S09]',
      'TMS：运输订单、计划配载、承运商委托、跟踪、自动计费、移动端、预约、车队/IoT、KPI。[S04][S08]',
      'AMS：仓库/月台容量、照单/无单/循环预约、签到、排队、靠台和作业分析。[S05]',
      '控制塔/AI：端到端事件、例外、BI/大屏、路径/装载/需求/库存/网络优化。[S06][S07]'])
    doc.add_heading('1.2 当前不能直接确认的内容',level=2)
    add_bullets(doc,['现网完整菜单树、全部字段、按钮级权限、隐藏特性开关、行业扩展和客户配置。','供应商真实源代码、数据库表、微服务边界、部署拓扑、第三方组件与私有接口。','当前登录后界面的像素级样式；公开手册中的界面可用于理解交互模式，但可能是历史版本。'])

    # 2
    add_section_title(doc,'2','能力地图与系统上下文')
    add_figure(doc,diagrams['01_system_context'],'图 2-1  系统上下文：内部用户、合作伙伴、企业系统和设备围绕统一供应链平台协同',17.0,10.5)
    add_figure(doc,diagrams['02_capability_map'],'图 2-2  能力地图：核心业务域、平台能力和可视化/优化能力',16.0,14.0)
    doc.add_heading('2.1 核心参与者',level=2)
    add_table(doc,['参与方','主要目标','典型入口'],[
      ['企业内部','订单履约、仓库作业、运输计划、预约现场、财务对账、运营决策','Web 工作台、RF/PDA、控制塔大屏'],
      ['客户/供应商','下单、确认、ASN、预约、跟踪、签收、评价、对账','客户/伙伴门户、APP、小程序'],
      ['承运商/司机','接单、报价、派车、节点上报、定位、异常、POD','承运商门户、司机 APP'],
      ['外部企业系统','同步主数据、订单、库存、财务、生产与电商数据','Open API、EDI、SFTP、Webhook'],
      ['设备/服务','扫描、打印、称重、门禁、定位、温控、地图、消息','IoT Hub、设备网关、第三方服务'],
    ],widths=[3.0,8.4,5.4],font_size=8.2)

    # 3 landscape, one figure per page where appropriate
    new_section(doc,True); add_section_title(doc,'3','建议系统架构','逻辑边界精确，部署形态可以从模块化单体平滑演进为服务化。')
    add_figure(doc,diagrams['03_logical_architecture'],'图 3-1  逻辑架构：渠道、编排、领域、平台基础能力和数据基础设施',25.6,12.2)
    add_body(doc,'领域服务不得通过跨库 JOIN 读取他域内部表。跨域写入使用命令/API 或领域事件；跨域查询由搜索和读模型聚合。')
    doc.add_page_break(); doc.add_heading('3.2 参考部署拓扑',level=2)
    add_figure(doc,diagrams['04_deployment'],'图 3-2  参考部署：边缘接入、无状态应用集群、多可用区数据层与运维安全',24.6,11.8)
    doc.add_page_break(); doc.add_heading('3.3 事件与数据架构',level=2)
    add_figure(doc,diagrams['05_event_data'],'图 3-3  事件与数据流：事务内 Outbox、消费者 Inbox、重试/死信和查询投影',13.0,12.1)
    doc.add_page_break(); doc.add_heading('3.4 多租户与安全架构',level=2)
    add_figure(doc,diagrams['06_tenancy_security'],'图 3-4  租户与安全：认证、RBAC+ABAC、数据隔离、脱敏、审计和密钥治理',13.0,12.1)
    doc.add_heading('3.5 架构原则',level=2)
    add_table(doc,['原则','落地要求'],[
      ['领域所有权','OMS、WMS、TMS、AMS、Billing、Control Tower 各自拥有状态机和数据库对象；外域只持 ID/必要快照。'],
      ['事件驱动','跨域传播版本化事实；Outbox/Inbox 解决提交与发布原子性、重复投递、重试和重放。'],
      ['规则可解释','分配、上架、波次、承运商、时隙和计费都保存规则版本、命中条件和排除原因。'],
      ['同步最小化','用户必须立即知道结果的命令同步；优化、批量、报表和跨域编排异步。'],
      ['配置版本化','流程、规则、模板、费率和字典通过发布版本生效，历史业务引用当时版本。'],
      ['可观测','traceId 与订单、任务、运单、预约、凭证业务号串联，可从一笔订单定位跨域步骤。'],
      ['渐进拆分','首期可模块化单体，但代码包、Schema、事件契约和团队责任按 bounded context 隔离。'],
    ],widths=[3.2,21.5],font_size=8.0)
    doc.add_heading('3.6 多租户数据隔离选择',level=2)
    add_table(doc,['模式','适用场景','优点','约束'],[
      ['共享库 + tenant_id','标准租户','成本低、运维简单','强制行级过滤、分区和防越权测试。'],
      ['共享实例 + 独立 Schema','中大型/监管租户','隔离清晰、迁移较易','Schema 数量和迁移编排复杂。'],
      ['独立数据库','超大租户/专属合规','性能、备份和加密隔离最佳','成本和版本运维更高。'],
      ['推荐：混合','默认共享，大租户可迁移','兼顾成本与隔离','ID、事件、配置和迁移工具从第一天支持。'],
    ],widths=[4.0,6.5,6.8,7.4],font_size=8.0)

    # 4 portrait
    new_section(doc,False); add_section_title(doc,'4','前端信息架构与页面规范','目标是高度相似的工作效率与信息密度，而不是复制品牌素材或受保护的前端代码。')
    add_figure(doc,diagrams['07_frontend_structure'],'图 4-1  前端结构：应用壳、核心页面模板与统一交互原语',16.0,10.8)
    add_callout(doc,'建议的视觉基线','可采用深色窄侧栏、浅色顶部上下文栏、多标签工作区、紧凑数据表格、右键/底部动作、右侧详情抽屉和统一状态色。品牌名、图标、插画、文案、CSS 与组件实现使用原创资产。','info')
    doc.add_heading('4.1 页面路由与模块清单',level=2); add_table(doc,['模块','建议路由','主要内容','模板'],UI_ROUTE_ROWS,widths=[3.0,4.1,7.0,2.7],font_size=7.7)
    doc.add_heading('4.2 前端组件契约',level=2); add_table(doc,['组件','行为与约束'],COMPONENT_ROWS,widths=[4.1,12.7],font_size=8.0)
    doc.add_heading('4.3 高相似交互细节',level=2)
    add_bullets(doc,[
      '表格默认紧凑密度：表头固定、服务端分页、行选中、首列业务号可进入详情；列配置、筛选和视图持久化到用户。',
      '动作同时可出现在顶部 CommandBar、行右键菜单或底部菜单，但必须由同一 Action Registry 决定权限、状态条件和确认文案。',
      '编辑页采用主信息+子表页签/侧边纵向页签；保存、审核、返回等关键命令固定位置，避免每个模块各自设计。',
      '异步动作立即返回 jobId 或受理状态，列表显示进度；成功 Toast 仅作短反馈，失败详情进入可复制的错误面板。',
      '键盘、扫码枪和移动端操作路径单独设计；RF 不直接复用桌面表格页面。'])

    # 5
    add_section_title(doc,'5','角色、权限与授权模型')
    add_table(doc,['角色','主要职责','关键边界'],ROLE_ROWS,widths=[3.3,7.4,6.1],font_size=7.8)
    doc.add_heading('5.1 命令授权顺序',level=2)
    add_bullets(doc,['认证：账号/令牌/设备是否有效，租户上下文是否可信。','资源授权：角色是否允许访问模块、API、动作和字段。','数据授权：组织树、仓库、货主、伙伴、创建者或自定义属性是否满足。','状态授权：当前业务状态是否允许动作，例如已发运不能直接取消。','业务规则：额度、证照、库存、费率、时间窗等是否满足。','审计：记录决策、操作者、前后值、IP/设备、traceId 和失败原因。'])

    # 6 features
    feature_count=sum(len(x[2]) for x in FEATURE_SECTIONS)
    add_section_title(doc,'6','详细功能规格',f'本节共 {feature_count} 项功能。A/B/C 表示证据等级，不等同于优先级。')
    for idx,(title,desc,rows) in enumerate(FEATURE_SECTIONS,1):
        h=doc.add_heading(f'6.{idx} {title}',level=2); keep_with_next(h,True); add_body(doc,desc)
        add_table(doc,['ID','功能与参与者','处理、规则与边界','输入/输出/状态','证据'],rows,widths=[1.65,3.25,7.15,3.55,1.15],font_size=6.95)

    # 7 flows portrait
    add_section_title(doc,'7','业务流程图')
    for i,(title,key,actors,pre,rules,exceptions) in enumerate(FLOW_DETAILS,1):
        if i>1: doc.add_page_break()
        doc.add_heading(f'7.{i} {title}',level=2)
        add_figure(doc,diagrams[key],f'图 7-{i}  {title}流程',15.5,16.3)
        add_table(doc,['项目','说明'],[['参与者','、'.join(actors)],['前置条件',pre],['关键规则',rules],['主要异常',exceptions]],widths=[2.8,14.0],font_size=8.0)

    # 8 sequences landscape
    new_section(doc,True); add_section_title(doc,'8','关键时序图')
    for i,(title,key,desc) in enumerate(SEQUENCE_DETAILS,1):
        if i>1: doc.add_page_break()
        doc.add_heading(f'8.{i} {title}',level=2)
        add_figure(doc,diagrams[key],f'图 8-{i}  {title}',25.0,11.8)
        add_body(doc,desc)

    # 9 states
    doc.add_page_break(); add_section_title(doc,'9','状态机与业务规则')
    state_keys=[('订单状态机','30_order_state'),('WMS 入库状态机','31_inbound_state'),('WMS 出库状态机','32_outbound_state'),('TMS 运输单状态机','33_shipment_state'),('预约状态机','34_appointment_state'),('结算凭证状态机','35_voucher_state')]
    for i,(title,key) in enumerate(state_keys,1):
        if i>1: doc.add_page_break()
        doc.add_heading(f'9.{i} {title}',level=2); add_figure(doc,diagrams[key],f'图 9-{i}  {title}',25.0,9.0)
    doc.add_page_break(); doc.add_heading('9.7 状态不变量摘要',level=2); add_table(doc,['对象','主路径','不可破坏规则'],STATE_RULE_ROWS,widths=[3.0,8.8,12.6],font_size=7.8)
    doc.add_heading('9.8 关键规则目录',level=2); add_table(doc,['规则域','主要输入','决策步骤','失败/解释'],RULE_ROWS,widths=[3.0,6.8,7.3,7.3],font_size=7.8)

    # 10 classes
    doc.add_page_break(); add_section_title(doc,'10','领域模型与类图','这些类图是拟建系统的 clean-room 领域模型，不声称是目标系统原始代码或数据库。')
    class_keys=[('领域聚合总览','40_class_overview'),('租户、权限与主数据','41_class_iam_master'),('OMS 订单与履约','42_class_oms'),('WMS 入库与库存','43_class_wms_inbound_inventory'),('WMS 出库与作业','44_class_wms_outbound'),('TMS 计划与执行','45_class_tms'),('AMS 预约与月台','46_class_ams'),('计费、控制塔、工作流与集成','47_class_billing_control')]
    for i,(title,key) in enumerate(class_keys,1):
        if i>1: doc.add_page_break()
        doc.add_heading(f'10.{i} {title}',level=2); add_figure(doc,diagrams[key],f'图 10-{i}  {title}类图',25.0,12.6)
    doc.add_page_break(); doc.add_heading('10.9 聚合边界与不变量',level=2); add_table(doc,['聚合','组成','核心不变量','数据所有权/事件'],AGGREGATE_ROWS,widths=[3.3,5.6,8.8,6.7],font_size=7.6)

    # 11 portrait
    new_section(doc,False); add_section_title(doc,'11','API、事件与一致性')
    doc.add_heading('11.1 API 约定',level=2)
    add_bullets(doc,['所有请求携带 Authorization、租户上下文和 X-Correlation-Id；外部写请求必须携带 Idempotency-Key。','命令返回对象 ID、版本和受理状态；长任务返回 jobId。查询使用服务端分页和稳定排序。','统一错误模型包含 code、message、fieldErrors、businessRef、correlationId、retryable；不向客户端返回堆栈。','附件采用预签名上传，上传完成后再提交元数据；API 版本以兼容性为导向。'])
    add_table(doc,['方法','路径','用途','关键语义'],API_ROWS,widths=[1.5,5.9,5.2,4.2],font_size=7.4)
    doc.add_heading('11.2 订单创建示例（建议接口，不是目标系统原接口）',level=2)
    add_code(doc,'''POST /api/v1/oms/orders
Authorization: Bearer <token>
X-Tenant-Id: tenant-demo
Idempotency-Key: erp-SO-20260713-0001-v1
X-Correlation-Id: 018f-order-trace

{
  "externalOrderNo": "SO-20260713-0001",
  "type": "SALES",
  "channel": "ERP",
  "customerId": "cus_123",
  "requestedDeliveryWindow": {
    "from": "2026-07-15T01:00:00Z",
    "to": "2026-07-15T09:00:00Z"
  },
  "lines": [{"lineNo": 1, "productId": "sku_001", "qty": 100, "uom": "EA"}]
}''')
    doc.add_heading('11.3 领域事件目录',level=2); add_table(doc,['事件','发布域','最小载荷','主要消费者'],EVENT_ROWS,widths=[3.3,2.3,7.2,4.0],font_size=7.2)
    doc.add_heading('11.4 一致性与失败语义',level=2)
    add_bullets(doc,['域内命令在单数据库事务内写业务表、状态变更、审计引用和 Outbox；发布器异步投递。','消费者先写 Inbox 去重，再执行业务；同一事件重复到达不得造成重复库存、重复任务或重复计费。','跨域长流程采用 Saga/过程管理器记录当前步骤、补偿命令和人工接管状态，不使用分布式事务。','事件按聚合键分区保持单对象顺序；跨对象不假设全局顺序，读模型使用事件版本忽略旧消息。','每日/期间运行订单-履约、库存-流水、运单-POD、计费-凭证对账任务，发现差异进入例外工单。'])

    # 12
    add_section_title(doc,'12','数据、安全与非功能')
    doc.add_heading('12.1 通用实体字段和建模规则',level=2)
    add_table(doc,['主题','规则'],[
      ['通用字段','id（UUID/有序 ID）、tenant_id、created_at/by、updated_at/by、version、status；业务流水和审计不软删。'],
      ['业务键','单号在 tenant+business_type 范围唯一；外部编号另建映射，不直接作为主键。'],
      ['金额','Money = amount + currency；税、汇率和舍入单独记录，禁止浮点数。'],
      ['数量','保存原单位数量和基础单位数量；包装换算引用当时的 PackageSpec 版本。'],
      ['地址/伙伴快照','业务单据保存主数据 ID 与关键快照，避免主数据后改改变历史语义。'],
      ['状态','状态变更只通过领域命令；更新携带 version；每次变更写领域事件和审计。'],
      ['扩展字段','使用受控 JSON Schema；可检索字段投影到索引，不用任意 JSON 代替核心结构。'],
      ['跨域引用','只保存对方 ID、业务号和必要快照，不建立跨域 ORM 导航或数据库外键。'],
    ],widths=[3.5,13.3],font_size=8.0)
    doc.add_heading('12.2 安全控制',level=2)
    add_bullets(doc,['认证：SSO/OIDC/SAML、本地账号、MFA、服务账号和短期令牌；高风险操作二次认证。','授权：RBAC 负责资源，ABAC 负责组织/仓库/货主/伙伴数据范围；后端统一策略决策。','数据：传输 TLS、静态 KMS 加密、对象存储私有桶、备份加密和密钥轮换。','隐私：字段脱敏、导出审批/水印、最小化收集、保留期和数据主体请求。','审计：登录、查询敏感信息、导出、变更、审批、接口和支持会话可追溯。'])
    doc.add_heading('12.3 非功能需求',level=2); add_table(doc,['主题','建议目标','设计措施'],NFR_ROWS,widths=[2.7,6.4,7.7],font_size=7.7)
    doc.add_heading('12.4 索引与数据分区建议',level=2)
    add_bullets(doc,['所有高频业务表索引以 tenant_id 开头；订单、任务、运单、预约和凭证的业务号建立租户范围唯一索引。','时间序列事件按 tenant+日期或业务域分区；轨迹/遥测进入专门时序或列式存储，在线库只留最近窗口。','库存余额唯一键覆盖 warehouse/location/owner/product/lot/serial/LPN/status；更新通过固定顺序或原子语句防死锁。','查询读模型按页面访问模式构建，不把在线事务库当跨域报表库。'])

    # 13
    add_section_title(doc,'13','实施路线与验收')
    add_table(doc,['阶段','范围','退出条件/交付物'],PHASE_ROWS,widths=[3.2,8.2,5.4],font_size=7.8)
    doc.add_heading('13.1 技术落地建议',level=2)
    add_bullets(doc,['前端：TypeScript + 成熟组件库；统一 Action Registry、状态词典、权限指令、表格和表单 schema；地图/甘特/日历独立模块。','后端：模块化单体或服务化均可，但每个领域独立包、Schema、仓储接口和事件；禁止跨模块直接访问内部表。','数据：关系数据库处理事务，Redis 用于缓存/短锁，消息总线处理事件，对象存储保存附件，搜索/OLAP 支撑跨域查询。','测试：状态机模型测试、库存/容量并发测试、幂等/重复/乱序事件测试、权限矩阵、合同费率回归样例和端到端对账。','发布：特性开关、灰度、可回滚数据库迁移、消息 schema 兼容和租户级配置发布。'])
    doc.add_heading('13.2 最小验收场景',level=2)
    add_bullets(doc,['ERP 创建销售订单→OMS 审核/分配→WMS 波次/拣选/包装/装车→TMS 发运/跟踪/POD→Billing 对账。','供应商 ASN→AMS 预约→门岗签到→WMS 收货/质检/上架→OMS/ERP 回写。','两名用户并发分配同一库存时不超卖；RF 重复/离线重放不重复扣减。','两名用户抢占同一预约时隙时仅一方成功；改期失败保持原预约。','承运商拒单后容量释放、运单回到可改派状态；POD 退回后不能结算。','费率变更不改变历史凭证；重算生成新版本并可解释差异。'])

    # 14
    add_section_title(doc,'14','现网核验清单','获得合法只读演示环境或脱敏录屏/HAR 后，用本节把 v0.1 升级为逐页、逐字段、逐动作的 v0.2。')
    add_table(doc,['核验主题','需要记录的内容'],VERIFY_ROWS,widths=[3.3,13.5],font_size=8.0)
    add_callout(doc,'建议的安全采集方式','使用临时只读演示账号，或提供脱敏菜单导出、全流程录屏、截图、网络 HAR 和样例打印件。不要共享生产管理员密码，也不要关闭访问控制。','warn')
    doc.add_heading('14.1 差异台账字段',level=2)
    add_table(doc,['字段','说明'],[
      ['页面/动作 ID','稳定编号，如 OMS-ORDER-LIST / ACTION-APPROVE。'],
      ['公开基线','本文对应功能 ID、流程、状态、类或 API。'],
      ['现网观察','页面、字段、默认值、权限、状态条件、请求/响应和视觉。'],
      ['差异类型','缺失、额外、规则差异、字段差异、视觉差异、角色差异、客户配置。'],
      ['决策','照现网、采用改进设计、兼容两种模式或放弃。'],
      ['证据','截图/录屏时间码/HAR 请求/测试账号角色。'],
    ],widths=[3.3,13.5],font_size=8.1)

    # 15 sources
    add_section_title(doc,'15','参考资料与证据说明')
    add_table(doc,['编号','资料','地址','用途与限制'],SOURCES,widths=[1.2,4.2,6.2,5.2],font_size=6.9)
    add_body(doc,'本文只抽象公开可观察的产品行为与供应链领域概念，没有复制公开手册中的截图、源代码或大段受版权保护文字。官方产品页用于确定能力边界；历史操作手册/演示材料用于理解流程和交互模式；工程架构、接口、数据模型和非功能要求属于建议实现。')

    path=BASE/'SCM_同类系统_架构与详细设计说明书_v0.1.docx'
    doc.save(path)
    return path


def build_mermaid() -> Path:
    text = '''# SCM 同类系统 - Mermaid 可编辑模型（v0.1）

> 这些图是 clean-room 建议模型，非目标站点原始代码或数据库。

## 1. 系统上下文
```mermaid
flowchart LR
  I[内部用户] --> P[供应链协同云平台]
  X[客户/供应商/承运商/司机] --> P
  E[ERP/MES/电商/财务] <--> P
  D[RF/打印/称重/门禁/GPS/IoT] <--> P
  P --> O[OMS]
  P --> W[WMS]
  P --> T[TMS]
  P --> A[AMS]
  P --> B[Billing]
  P --> C[Control Tower]
```

## 2. 端到端流程
```mermaid
flowchart TD
  A[订单接入] --> B[OMS 校验/审核/分配]
  B --> C[WMS 履约]
  B --> D[TMS 运输需求]
  C --> D
  D --> E[签收/POD]
  E --> F[计费/对账/财务]
  B -.异常.-> X[例外工单]
  C -.异常.-> X
  D -.异常.-> X
  X --> B
```

## 3. 订单状态机
```mermaid
stateDiagram-v2
  [*] --> Draft
  Draft --> Open: 提交
  Open --> Approved: 审核
  Approved --> Allocated: 分配
  Allocated --> Released: 释放
  Released --> Executing: 首个执行事件
  Executing --> Completed: 全部交付
  Completed --> Closed: 结算/归档
  Open --> Hold: 冻结
  Hold --> Open: 解冻
  Open --> Cancelled: 取消
```

## 4. 核心类图
```mermaid
classDiagram
  class BusinessOrder { UUID id; String orderNo; OrderStatus status; Long version }
  class OrderLine { UUID id; UUID productId; Decimal orderedQty }
  class FulfillmentOrder { UUID id; UUID warehouseId; FulfillmentStatus status }
  class InventoryBalance { UUID id; Decimal onHand; Decimal available; Long version }
  class Shipment { UUID id; String shipmentNo; ShipmentStatus status }
  class Appointment { UUID id; UUID slotId; AppointmentStatus status }
  class SettlementVoucher { UUID id; Money amount; VoucherStatus status }
  BusinessOrder "1" --> "*" OrderLine
  BusinessOrder "1" --> "*" FulfillmentOrder
  FulfillmentOrder ..> InventoryBalance
  BusinessOrder "1" --> "*" Shipment
  Shipment ..> Appointment
  Shipment ..> SettlementVoucher
```

## 5. 订单接入时序
```mermaid
sequenceDiagram
  participant ERP
  participant GW as API Gateway
  participant OMS
  participant BUS as Event Bus
  participant WMS
  participant TMS
  ERP->>GW: POST /orders + Idempotency-Key
  GW->>OMS: 鉴权/租户解析/命令
  OMS->>OMS: 校验、去重、状态机、Outbox
  OMS-->>ERP: 202 Accepted + orderId
  OMS-->>BUS: order.released.v1
  BUS-->>WMS: 创建履约任务
  BUS-->>TMS: 创建运输需求
```
'''
    p=BASE/'SCM_系统模型与流程图_Mermaid源文件_v0.1.md'; p.write_text(text,encoding='utf-8'); return p


def package_sources() -> Path:
    p=BASE/'SCM_架构流程类图_图源包_v0.1.zip'
    with zipfile.ZipFile(p,'w',zipfile.ZIP_DEFLATED) as z:
        for f in sorted(DIAG.iterdir()):
            if f.suffix.lower() in ('.dot','.svg','.png'): z.write(f,arcname=f'diagrams/{f.name}')
        mer=BASE/'SCM_系统模型与流程图_Mermaid源文件_v0.1.md'
        if mer.exists(): z.write(mer,arcname=mer.name)
    return p


def main():
    diagrams={}
    for fn in (architecture_diagrams,flow_diagrams,sequence_diagrams,state_diagrams,class_diagrams): diagrams.update(fn())
    docx=build_doc(diagrams)
    mer=build_mermaid()
    zipf=package_sources()
    print(json.dumps({'docx':str(docx),'mermaid':str(mer),'zip':str(zipf),'diagram_count':len(diagrams),'feature_count':sum(len(x[2]) for x in FEATURE_SECTIONS)},ensure_ascii=False,indent=2))

if __name__=='__main__':
    main()
